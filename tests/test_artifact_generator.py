from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

APP_ROOT = Path(__file__).resolve().parents[1] / "artpm_agent"

from artpm_agent.artifacts import WorkspaceArtifactGenerator
from artpm_agent.artifacts.verification import ArtifactVerificationError


def test_generate_xlsx_from_structured_rows_with_safe_metadata(tmp_path):
    root = tmp_path / "workspace" / "artifacts"
    generator = WorkspaceArtifactGenerator(root)

    result = generator.generate_xlsx(
        "项目报价.xlsx",
        {
            "sheet_name": "报价/明细",
            "columns": ["名称", "金额", "公式文本"],
            "rows": [
                {"名称": "角色", "金额": 1000, "公式文本": "=1+1"},
                ["场景", 2000, None],
            ],
        },
    )

    path = Path(result["path"])
    assert path.parent == root.resolve()
    assert result["name"] == "项目报价.xlsx"
    assert result["stored_path"] == result["name"]
    assert result["format"] == "xlsx"
    assert result["version"] == 1
    assert result["rows"] == 2
    assert result["columns"] == 3
    assert result["sheet_name"] == "报价_明细"
    assert result["size"] == path.stat().st_size
    assert result["sha256"] == sha256(path.read_bytes()).hexdigest()
    assert result["verification"]["status"] == "passed"
    assert result["verification"]["content_match"] is True
    assert result["verification"]["publication_integrity"] is True

    workbook = load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook["报价_明细"]
    values = list(worksheet.iter_rows(min_col=1, max_col=3, values_only=True))
    workbook.close()
    assert values == [
        ("名称", "金额", "公式文本"),
        ("角色", 1000, "'=1+1"),
        ("场景", 2000, None),
    ]


def test_generate_docx_from_structured_paragraphs(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    result = generator.generate_docx(
        "项目总结",
        [
            {"text": "交付总结", "kind": "heading", "level": 1},
            {"text": "本周按计划交付。", "bold": True},
            {"text": "检查源文件", "kind": "bullet", "italic": True},
            {"text": "通知客户", "kind": "numbered"},
        ],
    )

    path = Path(result["path"])
    assert result["name"] == "项目总结.docx"
    assert result["format"] == "docx"
    assert result["paragraphs"] == 4
    assert result["text_chars"] == sum(
        len(text)
        for text in ("交付总结", "本周按计划交付。", "检查源文件", "通知客户")
    )
    document = Document(path)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "交付总结",
        "本周按计划交付。",
        "检查源文件",
        "通知客户",
    ]
    assert document.paragraphs[0].style.name == "Title" or document.paragraphs[
        0
    ].style.name == "Heading 1"
    assert document.paragraphs[1].runs[0].bold is True
    assert document.paragraphs[2].style.name == "List Bullet"
    assert document.paragraphs[2].runs[0].italic is True
    assert document.paragraphs[3].style.name == "List Number"
    assert result["verification"]["status"] == "passed"
    assert result["verification"]["text_present"] is True


def test_generate_pptx_reopens_and_verifies_slide_text(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    result = generator.generate_pptx(
        "项目汇报.pptx",
        [
            {"title": "项目汇报", "bullets": ["范围已确认", "按计划交付"]},
            {"title": "下一步", "bullets": ["完成验收"]},
        ],
    )

    presentation = Presentation(result["path"])
    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert result["slides"] == 2
    assert result["verification"]["slide_count"] == 2
    assert result["verification"]["status"] == "passed"
    assert "范围已确认" in all_text
    assert "完成验收" in all_text


def test_generate_pdf_reopens_and_verifies_searchable_text(tmp_path):
    import pymupdf

    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    result = generator.generate_pdf(
        "项目简报.pdf",
        [
            {"text": "项目简报", "kind": "heading"},
            {"text": "你好，文件已经生成。"},
        ],
    )

    document = pymupdf.open(result["path"])
    try:
        extracted = "\n".join(page.get_text() for page in document)
        assert result["pages"] == document.page_count
    finally:
        document.close()
    assert result["verification"]["status"] == "passed"
    assert result["verification"]["text_present"] is True
    assert "你好，文件已经生成。" in extracted


def test_failed_file_verification_is_not_published(tmp_path, monkeypatch):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    def reject_output(*args, **kwargs):
        raise ArtifactVerificationError("corrupt output")

    monkeypatch.setattr(
        "artpm_agent.artifacts.generator.verify_artifact",
        reject_output,
    )

    with pytest.raises(ArtifactVerificationError, match="corrupt output"):
        generator.generate_docx("broken.docx", ["内容"])

    assert not list(generator.root.iterdir())


def test_filename_is_sanitized_and_collisions_create_versions(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    table = {"columns": ["值"], "rows": [[1]]}

    first = generator.generate_xlsx("季度:报告?.XLSX", table)
    first_bytes = Path(first["path"]).read_bytes()
    second = generator.generate_xlsx("季度:报告?.xlsx", table)

    assert first["name"] == "季度_报告_.xlsx"
    assert first["version"] == 1
    assert second["name"] == "季度_报告_ (2).xlsx"
    assert second["version"] == 2
    assert Path(first["path"]).read_bytes() == first_bytes
    assert not list((tmp_path / "artifacts").glob("*.tmp"))


def test_existing_same_name_is_never_overwritten(tmp_path):
    root = tmp_path / "artifacts"
    root.mkdir()
    existing = root / "report.xlsx"
    existing.write_bytes(b"existing-user-file")
    generator = WorkspaceArtifactGenerator(root)

    result = generator.generate_xlsx(
        "report.xlsx",
        {"columns": ["value"], "rows": [[1]]},
    )

    assert existing.read_bytes() == b"existing-user-file"
    assert result["name"] == "report (2).xlsx"
    assert Path(result["path"]).is_file()


@pytest.mark.parametrize(
    "filename",
    [
        "../escape.xlsx",
        "sub/escape.xlsx",
        "sub\\escape.xlsx",
        "C:\\escape.xlsx",
        "..",
    ],
)
def test_path_escape_and_path_components_are_rejected(tmp_path, filename):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    with pytest.raises(ValueError):
        generator.generate_xlsx(
            filename,
            {"columns": ["value"], "rows": []},
        )

    assert not list((tmp_path / "artifacts").iterdir())


def test_xlsx_row_column_cell_and_schema_limits(tmp_path):
    generator = WorkspaceArtifactGenerator(
        tmp_path / "artifacts",
        max_rows=1,
        max_columns=2,
        max_cell_chars=5,
    )

    with pytest.raises(ValueError, match="1 to 2 columns"):
        generator.generate_xlsx(
            "columns.xlsx",
            {"columns": ["a", "b", "c"], "rows": []},
        )
    with pytest.raises(ValueError, match="cannot exceed 1 rows"):
        generator.generate_xlsx(
            "rows.xlsx",
            {"columns": ["a"], "rows": [[1], [2]]},
        )
    with pytest.raises(ValueError, match="character limit"):
        generator.generate_xlsx(
            "cell.xlsx",
            {"columns": ["a"], "rows": [["123456"]]},
        )
    with pytest.raises(ValueError, match="unknown columns"):
        generator.generate_xlsx(
            "unknown.xlsx",
            {"columns": ["a"], "rows": [{"b": 1}]},
        )
    with pytest.raises(ValueError, match="must contain 2 cells"):
        generator.generate_xlsx(
            "width.xlsx",
            {"columns": ["a", "b"], "rows": [[1]]},
        )

    assert not list((tmp_path / "artifacts").iterdir())


def test_docx_paragraph_and_total_text_limits(tmp_path):
    generator = WorkspaceArtifactGenerator(
        tmp_path / "artifacts",
        max_paragraphs=1,
        max_total_text_chars=5,
    )

    with pytest.raises(ValueError, match="1 to 1 items"):
        generator.generate_docx("many.docx", ["a", "b"])
    with pytest.raises(ValueError, match="too large"):
        generator.generate_docx("large.docx", ["123456"])
    with pytest.raises(ValueError, match="unsupported fields"):
        generator.generate_docx("field.docx", [{"text": "a", "path": "x"}])

    assert not list((tmp_path / "artifacts").iterdir())


@pytest.mark.parametrize("method", ["generate_xlsx", "generate_docx"])
def test_generated_file_size_limit_removes_temporary_output(tmp_path, method):
    generator = WorkspaceArtifactGenerator(
        tmp_path / "artifacts",
        max_file_size=100,
    )

    with pytest.raises(ValueError, match="exceeds 100 bytes"):
        if method == "generate_xlsx":
            generator.generate_xlsx(
                "large.xlsx",
                {"columns": ["value"], "rows": [[1]]},
            )
        else:
            generator.generate_docx("large.docx", ["content"])

    assert not list((tmp_path / "artifacts").iterdir())


def test_concurrent_same_name_generation_is_atomic_and_versioned(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    table = {"columns": ["value"], "rows": [[1]]}

    with ThreadPoolExecutor(max_workers=8) as executor:
        results = list(
            executor.map(
                lambda _: generator.generate_xlsx("parallel.xlsx", table),
                range(8),
            )
        )

    assert len({result["name"] for result in results}) == 8
    assert {result["version"] for result in results} == set(range(1, 9))
    assert not list((tmp_path / "artifacts").glob("*.tmp"))
    for result in results:
        workbook = load_workbook(result["path"], read_only=True)
        assert next(workbook.active.iter_rows(values_only=True)) == ("value",)
        workbook.close()


def test_version_exhaustion_does_not_modify_existing_files(tmp_path):
    root = tmp_path / "artifacts"
    root.mkdir()
    (root / "report.docx").write_bytes(b"first")
    (root / "report (2).docx").write_bytes(b"second")
    generator = WorkspaceArtifactGenerator(root, max_versions=2)

    with pytest.raises(FileExistsError):
        generator.generate_docx("report.docx", ["new"])

    assert (root / "report.docx").read_bytes() == b"first"
    assert (root / "report (2).docx").read_bytes() == b"second"
    assert not list(root.glob("*.tmp"))


def test_mismatched_extension_and_non_directory_root_are_rejected(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    with pytest.raises(ValueError, match="extension must be .xlsx"):
        generator.generate_xlsx(
            "report.docx",
            {"columns": ["value"], "rows": []},
        )

    root_file = tmp_path / "not-a-directory"
    root_file.write_text("content", encoding="utf-8")
    with pytest.raises((FileExistsError, ValueError)):
        WorkspaceArtifactGenerator(root_file)


def test_generator_exposes_no_overwrite_or_delete_api(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    assert not hasattr(generator, "overwrite")
    assert not hasattr(generator, "delete")
    assert not hasattr(generator, "remove")


def test_generated_artifact_preview_and_export_formats(tmp_path):
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    artifact = generator.generate_xlsx(
        "quote.xlsx",
        {
            "columns": ["Item", "Cost"],
            "rows": [["Character", 1000], ["Scene", 2000]],
        },
    )

    assert "| Item | Cost |" in artifact["preview_markdown"]
    assert artifact["export_formats"] == ["csv", "md", "txt", "docx"]

    preview = generator.preview_artifact(artifact["stored_path"])
    assert preview["success"] is True
    assert "Character" in preview["preview_markdown"]

    csv_export = generator.export_artifact_bytes(artifact["stored_path"], "csv")
    assert csv_export["filename"] == "quote.csv"
    assert "Item,Cost" in csv_export["data"].decode("utf-8-sig")

    md_export = generator.export_artifact_bytes(artifact["stored_path"], "md")
    assert b"| Item | Cost |" in md_export["data"]

    docx_export = generator.export_artifact_bytes(artifact["stored_path"], "docx")
    assert docx_export["filename"] == "quote.docx"
    assert docx_export["data"].startswith(b"PK")

    saved = generator.export_artifact(artifact["stored_path"], "md")
    assert saved["format"] == "md"
    assert saved["source_artifact"] == artifact["stored_path"]
    assert Path(saved["path"]).is_file()
