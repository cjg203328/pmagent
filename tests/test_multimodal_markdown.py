from __future__ import annotations

from docx import Document
from openpyxl import Workbook
from PIL import Image
import pytest

from artpm_agent.agent import ArtPMAgent
from artpm_agent.harness.knowledge_handler import _build_knowledge_ingestion_resources
from artpm_agent.utils.multimodal_markdown import LocalMarkdownConverter


def _write_workbook(path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Quote"
    worksheet.append(["Asset", "Qty", "Cost"])
    worksheet.append(["Character Model", 2, 1200])
    worksheet.append(["Scene Prop", 1, 300])
    workbook.save(path)


def test_converter_normalizes_spreadsheet_to_markdown(tmp_path):
    xlsx_path = tmp_path / "quote.xlsx"
    _write_workbook(xlsx_path)

    result = LocalMarkdownConverter().convert(xlsx_path)

    assert result.success is True
    assert result.source_format == ".xlsx"
    assert "## Sheet: Quote" in result.markdown
    assert "| Asset | Qty | Cost |" in result.markdown
    assert "| Character Model | 2 | 1200 |" in result.markdown


def test_converter_normalizes_docx_and_image_ocr_text(tmp_path):
    docx_path = tmp_path / "brief.docx"
    document = Document()
    document.add_heading("Production Brief", level=1)
    document.add_paragraph("Use the approved asset naming rules.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Owner"
    table.cell(1, 1).text = "ArtPM"
    document.save(docx_path)

    converter = LocalMarkdownConverter()
    docx_result = converter.convert(docx_path)

    assert docx_result.success is True
    assert "## Production Brief" in docx_result.markdown
    assert "| Field | Value |" in docx_result.markdown
    assert "| Owner | ArtPM |" in docx_result.markdown

    image_path = tmp_path / "capture.png"
    Image.new("RGB", (16, 12), "white").save(image_path, "PNG")
    image_result = converter.convert(
        image_path,
        parsed_result={
            "raw_text": "Detected table text",
            "extracted_data": {"width": 16, "height": 12},
            "requires_vision": False,
            "ocr_available": True,
            "ocr_status": "completed",
        },
    )

    assert image_result.success is True
    assert "## OCR Text" in image_result.markdown
    assert "Detected table text" in image_result.markdown
    assert image_result.ocr_status == "completed"


def test_agent_process_document_adds_markdown_preprocessor(tmp_path):
    xlsx_path = tmp_path / "quote.xlsx"
    _write_workbook(xlsx_path)

    class Router:
        def execute_skill(self, skill_name, inputs):
            assert skill_name == "document_classifier_parser"
            assert inputs["file_path"] == str(xlsx_path)
            return {
                "success": True,
                "document_type": "Excel",
                "extracted_data": {"parser": "legacy"},
                "raw_text": "",
            }

    agent = object.__new__(ArtPMAgent)
    agent.router = Router()
    agent.markdown_converter = LocalMarkdownConverter()

    result = agent.process_document(str(xlsx_path), "learn this table style")

    assert result["markdown"].startswith("# quote.xlsx")
    assert result["raw_text"] == result["markdown"]
    assert result["preprocessor"]["name"] == "local-markdown"
    assert result["extracted_data"]["markdown"]["source_format"] == ".xlsx"
    assert "| Asset | Qty | Cost |" in result["markdown"]


def test_agent_process_document_does_not_index_empty_image_ocr(tmp_path):
    image_path = tmp_path / "capture.png"
    Image.new("RGB", (16, 12), "white").save(image_path, "PNG")

    class Router:
        def execute_skill(self, skill_name, inputs):
            return {
                "success": True,
                "document_type": "Image",
                "extracted_data": {"width": 16, "height": 12},
                "raw_text": "",
                "requires_vision": True,
                "ocr_status": "unavailable",
            }

    agent = object.__new__(ArtPMAgent)
    agent.router = Router()
    agent.markdown_converter = LocalMarkdownConverter()

    result = agent.process_document(str(image_path), "read image")

    assert result["raw_text"] == ""
    assert "No OCR text extracted." in result["markdown"]
    assert result["preprocessor"]["source_format"] == ".png"


def test_attachment_context_includes_markdown_evidence(tmp_path):
    file_path = tmp_path / "quote.xlsx"
    file_path.write_bytes(b"placeholder")

    agent = object.__new__(ArtPMAgent)
    agent.process_document = lambda path, hint: {
        "success": True,
        "document_type": "Excel",
        "extracted_data": {"rows": 2},
        "raw_text": "# quote.xlsx\n\nraw fallback",
        "markdown": "# quote.xlsx\n\n| Asset | Qty |\n| --- | --- |\n| A | 1 |",
        "preprocessor": {"name": "local-markdown"},
    }

    parsed, context = agent._parse_context_attachments(
        "build the same table",
        {"file_paths": [str(file_path)]},
    )

    assert parsed[0]["markdown"].startswith("# quote.xlsx")
    assert "<attachment_markdown>" in context
    assert "| Asset | Qty |" in context
    assert context.count("<attachment_data>") == 1


def test_knowledge_ingestion_uses_markdown_without_empty_image_noise():
    class Agent:
        def __init__(self, result):
            self.result = result

        def process_document(self, _path, _prompt):
            return self.result

    resources = _build_knowledge_ingestion_resources(
        Agent(
            {
                "success": True,
                "document_type": "Excel",
                "extracted_data": {"rows": 2},
                "raw_text": "",
                "markdown": "# quote.xlsx\n\n| Asset | Qty |\n| --- | --- |\n| A | 1 |",
                "preprocessor": {"name": "local-markdown"},
            }
        ),
        [{"name": "quote.xlsx", "extension": "xlsx"}],
        ["quote.xlsx"],
        "learn this file",
    )

    assert resources[0]["resource_type"] == "table"
    assert resources[0]["searchable_text"].startswith("# quote.xlsx")
    assert resources[0]["metadata"]["preprocessor"]["name"] == "local-markdown"

    with pytest.raises(ValueError):
        _build_knowledge_ingestion_resources(
            Agent(
                {
                    "success": True,
                    "document_type": "Image",
                    "extracted_data": {"width": 16},
                    "raw_text": "",
                    "markdown": "# capture.png\n\nNo OCR text extracted.",
                    "preprocessor": {"name": "local-markdown"},
                }
            ),
            [{"name": "capture.png", "extension": "png"}],
            ["capture.png"],
            "learn this image",
        )
