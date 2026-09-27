"""S0 回归：Excel 分类判别、按月份裁剪、XLSX 生成与核验。"""

from __future__ import annotations

import io
import re
import shutil
import zipfile
from hashlib import sha256
from pathlib import Path

import openpyxl
import pytest
from openpyxl import load_workbook

from artpm_agent.artifacts import WorkspaceArtifactGenerator
from artpm_agent.artifacts.verification import (
    ArtifactVerificationError,
    verify_artifact,
)
from artpm_agent.artifacts.xlsx_monthly import (
    MonthlyXlsxError,
    normalize_cell_month,
    normalize_target_month,
    transform_monthly_workbook,
)
from artpm_agent.parsers.excel_parser import ExcelQuoteParser


def _stream(rows, name="probe.xlsx") -> io.BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    stream.name = name
    workbook.save(stream)
    stream.seek(0)
    return stream


def _monthly_source(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "人天分配"
    sheet.append(["任务", "月份", "商务人天", "已分配人天"])
    sheet.append(["麦迪", "2026-10", 18, 18])
    sheet.append(["麦迪", "2026-11", 18, 0])
    sheet.append(["裸模优化", "2026-11", 15, 0])
    sheet.append(["场景贴图", "2026-12", 12, 0])
    sheet.column_dimensions["A"].width = 22
    workbook.save(path)
    workbook.close()
    return path


# ── R-3 分类判别 ──────────────────────────────────────────────


def test_manual_allocation_sheet_is_not_classified_as_quote():
    """人天表不得被判成报价单（PRD §10 R-3）。"""
    result = ExcelQuoteParser().parse(
        _stream(
            [
                ["角色组-任务及绩效分配"],
                ["任务", "商务人天", "已分配人天", "差异", "状态"],
                ["麦迪", 18, 0, -18, "未开始分配"],
            ]
        )
    )
    assert result["success"] is True
    assert result["document_type"] in {"人天分配表", "产能表", "unknown"}
    assert result["document_type"] != "报价单"
    assert result["assets"] == []


def test_unrecognizable_sheet_returns_unknown():
    result = ExcelQuoteParser().parse(_stream([["foo", "bar"], ["a", "b"]]))
    assert result["document_type"] == "unknown"


def test_quote_sheet_with_unknown_client_stays_a_quote():
    """客户识别失败不得连带把报价表降级为 unknown。"""
    result = ExcelQuoteParser().parse(
        _stream([["资产名称", "数量", "单价", "总价"], ["主角", 2, 1000, 2000]])
    )
    assert result["document_type"] == "报价单"
    assert result["total_amount"] == 2000
    assert len(result["assets"]) == 1


# ── 月份归一化 ────────────────────────────────────────────────


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("2026-11", "2026-11"),
        ("2026年11月", "2026-11"),
        ("2026/11", "2026-11"),
        ("2026.11", "2026-11"),
    ],
)
def test_normalize_target_month_accepts_common_shapes(raw, expected):
    assert normalize_target_month(raw) == expected


def test_normalize_target_month_rejects_garbage():
    with pytest.raises(MonthlyXlsxError):
        normalize_target_month("下个月")


def test_normalize_cell_month_handles_date_and_month_only():
    import datetime

    assert normalize_cell_month(datetime.date(2026, 11, 20)) == "2026-11"
    assert normalize_cell_month("11月", default_year=2026) == "2026-11"
    assert normalize_cell_month("", default_year=2026) is None
    assert normalize_cell_month(202611, default_year=2026) is None


# ── 裁剪行为 ──────────────────────────────────────────────────


def test_transform_keeps_only_target_month_rows(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")
    output = tmp_path / "november.xlsx"

    result = transform_monthly_workbook(source, output, "2026-11")

    assert result.target_month == "2026-11"
    assert result.mode == "new_items"
    assert result.matched_rows == 2
    assert result.uncovered_rows == 2
    assert output.is_file()

    workbook = load_workbook(output)
    sheet = workbook["人天分配"]
    kept = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
    headers = [sheet.cell(1, col).value for col in range(1, 5)]
    workbook.close()

    assert kept == ["麦迪", "裸模优化"]
    assert headers == ["任务", "月份", "商务人天", "已分配人天"]


def test_transform_preserves_column_width(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")
    output = tmp_path / "november.xlsx"

    transform_monthly_workbook(source, output, "2026-11")

    workbook = load_workbook(output)
    width = workbook["人天分配"].column_dimensions["A"].width
    workbook.close()
    assert width == pytest.approx(22)


def test_cumulative_mode_keeps_earlier_months(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")
    output = tmp_path / "cumulative.xlsx"

    result = transform_monthly_workbook(source, output, "2026-11", mode="cumulative")

    assert result.mode == "cumulative"
    assert result.matched_rows == 3

    workbook = load_workbook(output)
    kept = [
        workbook["人天分配"].cell(row, 1).value
        for row in range(2, workbook["人天分配"].max_row + 1)
    ]
    workbook.close()
    assert kept == ["麦迪", "麦迪", "裸模优化"]


def test_transform_rejects_legacy_xls(tmp_path):
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(MonthlyXlsxError):
        transform_monthly_workbook(source, tmp_path / "out.xlsx", "2026-11")


def test_transform_rejects_unknown_mode(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")

    with pytest.raises(MonthlyXlsxError):
        transform_monthly_workbook(source, tmp_path / "out.xlsx", "2026-11", mode="all")


def test_transform_raises_when_no_month_column(tmp_path):
    source = tmp_path / "no_month.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["名称", "数量"])
    sheet.append(["角色", 3])
    workbook.save(source)
    workbook.close()

    with pytest.raises(MonthlyXlsxError):
        transform_monthly_workbook(source, tmp_path / "out.xlsx", "2026-11")


# ── 口径说明 ──────────────────────────────────────────────────


def test_processing_note_records_mode_and_counts(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")
    output = tmp_path / "november.xlsx"

    transform_monthly_workbook(source, output, "2026-11")

    workbook = load_workbook(output)
    notes = {
        row[0].value: row[1].value
        for row in workbook["处理说明"].iter_rows(min_row=1, values_only=False)
        if row and row[0].value is not None
    }
    workbook.close()

    assert notes["目标月份"] == "2026-11"
    assert notes["处理口径"] == "目标月新增"
    assert int(notes["保留行数"]) == 2
    assert int(notes["未覆盖行数"]) == 2
    assert notes["源文件 SHA-256"] == sha256(source.read_bytes()).hexdigest()


# ── 生成 + 核验 ───────────────────────────────────────────────


def test_generator_monthly_artifact_passes_verification(tmp_path):
    root = tmp_path / "workspace"
    source = _monthly_source(tmp_path / "source.xlsx")
    generator = WorkspaceArtifactGenerator(root / "artifacts")

    result = generator.generate_xlsx_monthly("11月汇总.xlsx", source, "2026-11")

    path = Path(result["path"])
    assert path.is_file()
    assert result["verification"]["status"] == "passed"
    assert result["verification"]["reopened"] is True
    assert result["verification"]["monthly_target"] == "2026-11"
    assert (
        result["verification"]["source_sha256"]
        == sha256(source.read_bytes()).hexdigest()
    )
    assert result["sha256"] == sha256(path.read_bytes()).hexdigest()


def test_verification_rejects_tampered_monthly_workbook(tmp_path):
    source = _monthly_source(tmp_path / "source.xlsx")
    output = tmp_path / "november.xlsx"
    transform_monthly_workbook(source, output, "2026-11")

    with pytest.raises(ArtifactVerificationError):
        verify_artifact(
            output,
            "xlsx_monthly",
            {
                "target_month": "2026-12",
                "mode": "new_items",
                "source_sha256": sha256(source.read_bytes()).hexdigest(),
                "matched_rows": 2,
            },
        )


def test_verification_rejects_plain_workbook_as_monthly(tmp_path):
    """普通工作簿缺少处理说明表，不得被当作月度产物放行。"""
    plain = tmp_path / "plain.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["名称", "数量"])
    workbook.save(plain)
    workbook.close()

    with pytest.raises(ArtifactVerificationError):
        verify_artifact(plain, "xlsx_monthly", {"target_month": "2026-11"})


def test_monthly_run_is_deterministic(tmp_path):
    """同一输入两次运行，保留行与口径说明必须一致。"""
    source = _monthly_source(tmp_path / "source.xlsx")
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"

    one = transform_monthly_workbook(source, first, "2026-11")
    two = transform_monthly_workbook(source, second, "2026-11")

    assert one.to_dict()["note"] == two.to_dict()["note"]
    assert one.matched_rows == two.matched_rows
    assert one.sheets == two.sheets

    def rows_of(path):
        workbook = load_workbook(path)
        values = [
            tuple(row) for row in workbook["人天分配"].iter_rows(values_only=True)
        ]
        workbook.close()
        return values

    assert rows_of(first) == rows_of(second)


# ── 宽表：月份在 sheet 名上（真实附件 `角色组-任务及绩效分配.xlsx` 的形状）──


def _wide_source(path: Path, titles=("24年10月", "24年11月", "25年1月")) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for index, title in enumerate(titles):
        sheet = workbook.create_sheet(title)
        sheet.append(["甲方工作室", "任务名称", "商务人天", "已分配人天"])
        sheet.append(["NBA", "麦迪", 18, index])
        sheet.append(["", "艾弗森", 19, index + 1])
    workbook.save(path)
    return path


@pytest.mark.parametrize(
    ("title", "expected"),
    [("24年11月", "2024-11"), ("2024-11", "2024-11"), ("25年1月", "2025-01")],
)
def test_sheet_title_is_read_as_a_month(title, expected):
    from artpm_agent.artifacts.xlsx_monthly import sheet_target_month

    assert sheet_target_month(title) == expected


@pytest.mark.parametrize("title", ["人天分配", "处理说明", "24年", "汇总"])
def test_non_month_sheet_title_is_not_a_month(title):
    from artpm_agent.artifacts.xlsx_monthly import sheet_target_month

    assert sheet_target_month(title) is None


def test_wide_workbook_selects_the_matching_sheet(tmp_path):
    source = _wide_source(tmp_path / "wide.xlsx")
    output = tmp_path / "out.xlsx"

    result = transform_monthly_workbook(source, output, "2024-11")

    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["24年11月", "处理说明"]
    finally:
        workbook.close()
    assert result.matched_rows == 2
    assert result.month_column == "工作表名称"


def test_wide_workbook_cumulative_keeps_earlier_months(tmp_path):
    source = _wide_source(tmp_path / "wide.xlsx")
    output = tmp_path / "out.xlsx"

    transform_monthly_workbook(source, output, "2024-11", mode="cumulative")

    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["24年10月", "24年11月", "处理说明"]
    finally:
        workbook.close()


def test_bare_month_resolves_from_sheet_titles_without_asking(tmp_path):
    """「仅11月汇总版」不带年份，但工作簿只用过 2024 年 11 月，因此无歧义。"""
    from artpm_agent.artifacts.xlsx_monthly import resolve_target_month

    source = _wide_source(tmp_path / "wide.xlsx")

    assert resolve_target_month("要「仅11月汇总版」", source) == "2024-11"


def test_ambiguous_month_across_years_asks_instead_of_guessing(tmp_path):
    from artpm_agent.artifacts.xlsx_monthly import (
        MonthlyXlsxError,
        resolve_target_month,
    )

    source = _wide_source(tmp_path / "wide.xlsx", titles=["24年3月", "25年3月"])

    with pytest.raises(MonthlyXlsxError, match="请指定年份"):
        resolve_target_month("保留3月", source)


@pytest.mark.parametrize(
    ("prompt", "expected"),
    [("只保留24年11月的行", "2024-11"), ("只保留25年11月的行", "2025-11")],
)
def test_two_digit_year_prompt_is_anchored_on_the_workbook(tmp_path, prompt, expected):
    """表名写作「25年11月」，用户也会这样提需求；年份不得被当成歧义拒掉。"""
    from artpm_agent.artifacts.xlsx_monthly import resolve_target_month

    source = _wide_source(tmp_path / "wide.xlsx", titles=["24年11月", "25年11月"])

    assert resolve_target_month(prompt, source) == expected


def test_two_digit_year_without_a_workbook_still_asks():
    from artpm_agent.artifacts.xlsx_monthly import (
        MonthlyXlsxError,
        resolve_target_month,
    )

    with pytest.raises(MonthlyXlsxError, match="四位年份"):
        resolve_target_month("只保留25年11月的行")


def _with_empty_fill(path: Path) -> Path:
    """Add a self-closing ``<fill/>`` node, as WPS and some add-ins export."""
    with zipfile.ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    styles = {item.filename: payload for item, payload in members}[
        "xl/styles.xml"
    ].decode("utf-8")
    styles, updated = re.subn(
        r'<fills count="(\d+)">',
        lambda match: f'<fills count="{int(match.group(1)) + 1}">',
        styles,
        count=1,
    )
    assert updated == 1, "openpyxl output did not contain a countable <fills> node"
    styles = styles.replace("</fills>", "<fill/></fills>", 1).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, payload in members:
            archive.writestr(
                item, styles if item.filename == "xl/styles.xml" else payload
            )
    return path


def test_wide_row_count_ignores_formatting_only_tail_rows(tmp_path):
    """真实月度表的格式延伸到最后一行之后，空行不得算成交付行数。"""
    from openpyxl.styles import PatternFill

    source = tmp_path / "wide.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("24年11月")
    sheet.append(["甲方工作室", "任务名称", "商务人天"])
    sheet.append(["NBA", "麦迪", 18])
    sheet.append(["", "艾弗森", 19])
    sheet.cell(200, 3).fill = PatternFill(
        fill_type="solid", start_color="FF8CDDFA", end_color="FF8CDDFA"
    )
    workbook.save(source)

    output = tmp_path / "out.xlsx"
    result = transform_monthly_workbook(source, output, "2024-11")

    assert sheet.max_row == 200, "夹具必须复现真实表的格式延伸"
    assert result.matched_rows == 2


def test_monthly_transform_opens_a_workbook_with_empty_fill_nodes(tmp_path):
    """真实客户表用 WPS 导出，空 ``<fill/>`` 会让 openpyxl 整本报错（PRD §10）。"""
    source = _with_empty_fill(_wide_source(tmp_path / "wide.xlsx"))

    with pytest.raises(TypeError):
        load_workbook(source)

    output = tmp_path / "out.xlsx"
    result = transform_monthly_workbook(source, output, "2024-11")

    assert result.matched_rows == 2
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["24年11月", "处理说明"]
    finally:
        workbook.close()


# ── 端到端：请求必须经 run_turn 走到可下载交付物，而不是聊天内表格 ──


class _NoLLM:
    def chat(self, *_args, **_kwargs):
        raise AssertionError("月度裁剪是确定性路径，不应调用模型")


class _Capabilities:
    # 月度裁剪只读原始工作簿，不需要解析器参与本轮。
    attachment_parsing = False


class _Runtime:
    capabilities = _Capabilities()


def _run_turn_monthly(tmp_path, prompt):
    from artpm_agent.artifacts.coordinator import ArtifactCoordinator
    from artpm_agent.harness import run_turn
    from artpm_agent.harness.turn_service import TurnContext

    source = _wide_source(tmp_path / "角色组-任务及绩效分配.xlsx")
    root = tmp_path / "artifacts"
    root.mkdir()
    coordinator = ArtifactCoordinator(
        WorkspaceArtifactGenerator(root),
        llm=_NoLLM(),
    )
    context = TurnContext(
        turn_id="turn-s0",
        conversation_id="conv-s0",
        user_input=prompt,
        attachments=[
            {
                "name": source.name,
                "extension": "xlsx",
                "stored_path": str(source),
            }
        ],
        extra={"file_paths": [str(source)]},
        runtime=_Runtime(),
    )
    result = run_turn(
        context,
        artifact_coordinator=coordinator,
        request_conversation_id="conv-s0",
    )
    return result, root


@pytest.mark.parametrize(
    "prompt",
    [
        "要「仅11月汇总版」",
        "帮我把表格数据 去除掉 只保留完整的11月",
    ],
)
def test_run_turn_delivers_a_monthly_workbook(tmp_path, prompt):
    """S0 的验收判据：产出文件，而不是把表格打印在回复里。"""
    result, root = _run_turn_monthly(tmp_path, prompt)

    assert result.success is True
    assert result.handled_by == "artifact_generation"
    assert result.artifacts, "月度请求必须产出交付物"
    artifact = result.artifacts[0]
    assert artifact["verification"]["status"] == "passed"
    delivered = root / artifact["stored_path"]
    assert delivered.is_file()

    workbook = load_workbook(delivered)
    try:
        assert workbook.sheetnames == ["24年11月", "处理说明"]
    finally:
        workbook.close()


def test_monthly_artifact_is_named_after_the_uploaded_file_and_month(tmp_path):
    """上传件在盘上是哈希名，交付物必须回到用户认得的名字（含月份）。"""
    from artpm_agent.artifacts.coordinator import ArtifactCoordinator

    stored = tmp_path / "uploads" / "58af653edd434cf68f0b661dd780bbf4.xlsx"
    stored.parent.mkdir()
    shutil.copy(_wide_source(tmp_path / "wide.xlsx"), stored)
    root = tmp_path / "artifacts"
    root.mkdir()
    coordinator = ArtifactCoordinator(
        WorkspaceArtifactGenerator(root),
        llm=_NoLLM(),
    )

    result = coordinator.process(
        "只保留24年11月的行",
        attachments=[
            {"name": "角色组-任务及绩效分配.xlsx", "stored_path": str(stored)}
        ],
        file_paths=[str(stored)],
    )

    assert result.rejected is False
    assert result.artifact["name"] == "角色组-任务及绩效分配-2024-11.xlsx"
