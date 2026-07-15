import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from artpm_agent.artifacts import (
    ArtifactCoordinator,
    SpreadsheetTemplateStore,
    WorkspaceArtifactGenerator,
)


class FakeLLM:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def chat(self, prompt, **kwargs):
        self.calls.append({"prompt": prompt, **kwargs})
        if isinstance(self.response, Exception):
            raise self.response
        return self.response


def write_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cost"
    sheet.append(["Item", "Qty", "Unit"])
    sheet.append(["Role", 2, 100])
    workbook.save(path)
    workbook.close()


def test_template_store_learns_searches_and_records_use(tmp_path):
    source = tmp_path / "source.xlsx"
    write_template(source)
    store = SpreadsheetTemplateStore(tmp_path / "templates.json")

    template = store.learn_from_file(
        source,
        name="Cost Template",
        keywords=["quote"],
        source={"name": "source.xlsx"},
    )

    assert template["name"] == "Cost Template"
    assert template["sheets"][0]["name"] == "Cost"
    assert template["sheets"][0]["columns"] == ["Item", "Qty", "Unit"]
    assert template["sheets"][0]["sample_rows"][0]["Item"] == "Role"

    reopened = SpreadsheetTemplateStore(tmp_path / "templates.json")
    matches = reopened.search("generate quote using Cost Template")
    assert matches[0]["id"] == template["id"]

    updated = reopened.record_use(template["id"])
    assert updated["usage_count"] == 1


def test_generator_preserves_template_sheet_columns_and_versions(tmp_path):
    source = tmp_path / "source.xlsx"
    write_template(source)
    template = SpreadsheetTemplateStore(tmp_path / "templates.json").learn_from_file(
        source,
        name="Cost Template",
    )
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    result = generator.generate_xlsx_from_template(
        "cost.xlsx",
        template,
        rows_by_sheet={"Cost": [["Scene", 1, 200]]},
    )

    workbook = load_workbook(result["path"], read_only=True)
    sheet = workbook["Cost"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    assert rows == [("Item", "Qty", "Unit"), ("Scene", 1, 200)]
    assert result["rows"] == 1
    assert result["columns"] == 3
    assert result["template_name"] == "Cost Template"


def test_coordinator_learns_template_then_generates_from_it_offline(tmp_path):
    source = tmp_path / "source.xlsx"
    write_template(source)
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    template_store = SpreadsheetTemplateStore(tmp_path / "templates.json")
    llm = FakeLLM(RuntimeError("offline"))
    coordinator = ArtifactCoordinator(generator, llm, template_store)

    learned = coordinator.process(
        "learn this table format as template Cost Template",
        attachments=[{"name": "source.xlsx", "sha256": "sha"}],
        file_paths=[source],
    )

    assert learned.matched is True
    assert learned.error_code is None
    assert learned.artifact is None
    assert "Cost Template" in learned.message
    assert llm.calls == []

    generated = coordinator.process(
        "generate xlsx using template Cost Template with data: Scene, 1, 200"
    )

    assert generated.matched is True
    assert generated.error_code is None
    assert generated.artifact["format"] == "xlsx"
    assert generated.artifact["rows"] == 1

    workbook = load_workbook(generated.artifact["path"], read_only=True)
    rows = list(workbook["Cost"].iter_rows(values_only=True))
    workbook.close()
    assert rows == [("Item", "Qty", "Unit"), ("Scene", "1", "200")]
    assert len(llm.calls) == 1


def test_coordinator_constrains_llm_rows_to_template_columns(tmp_path):
    source = tmp_path / "source.xlsx"
    write_template(source)
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    template_store = SpreadsheetTemplateStore(tmp_path / "templates.json")
    template_store.learn_from_file(source, name="Cost Template")
    llm = FakeLLM(
        json.dumps(
            {
                "format": "xlsx",
                "filename": "custom.xlsx",
                "table": {
                    "sheet_name": "Any",
                    "columns": ["Unit", "Item", "Qty"],
                    "rows": [[300, "Prop", 4]],
                },
            }
        )
    )
    coordinator = ArtifactCoordinator(generator, llm, template_store)

    generated = coordinator.process(
        "generate excel using template Cost Template with data: ignored"
    )

    workbook = load_workbook(generated.artifact["path"], read_only=True)
    rows = list(workbook["Cost"].iter_rows(values_only=True))
    workbook.close()

    assert generated.artifact["name"] == "custom.xlsx"
    assert rows == [("Item", "Qty", "Unit"), ("Prop", 4, 300)]
    assert "Cost Template" in llm.calls[0]["system_prompt"]
