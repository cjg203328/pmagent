import json
from pathlib import Path

from docx import Document

from artpm_agent.artifacts import (
    ArtifactCoordinator,
    DocumentTemplateStore,
    WorkspaceArtifactGenerator,
)


class FakeLLM:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def chat(self, prompt, **kwargs):
        self.calls.append({"prompt": prompt, **kwargs})
        if isinstance(self.response, Exception):
            raise self.response
        return self.response


def write_doc_template(path: Path) -> None:
    document = Document()
    document.add_heading("Weekly Report", level=1)
    paragraph = document.add_paragraph("Summary")
    paragraph.runs[0].bold = True
    document.add_paragraph("Milestone", style="List Bullet")
    document.save(path)


def test_document_template_store_learns_outline_and_searches(tmp_path):
    source = tmp_path / "weekly.docx"
    write_doc_template(source)
    store = DocumentTemplateStore(tmp_path / "doc_templates.json")

    template = store.learn_from_file(source, name="Weekly Template")

    assert template["kind"] == "document"
    assert template["name"] == "Weekly Template"
    assert [item["kind"] for item in template["paragraphs"]] == [
        "heading",
        "paragraph",
        "bullet",
    ]
    assert template["paragraphs"][0]["level"] == 1
    assert template["paragraphs"][1]["bold"] is True
    assert store.search("generate Weekly Template doc")[0]["id"] == template["id"]


def test_generator_preserves_document_template_structure(tmp_path):
    source = tmp_path / "weekly.docx"
    write_doc_template(source)
    template = DocumentTemplateStore(tmp_path / "doc_templates.json").learn_from_file(
        source,
        name="Weekly Template",
    )
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")

    result = generator.generate_docx_from_template(
        "report.docx",
        template,
        paragraphs=[
            {"text": "Production Update"},
            {"text": "Revenue and delivery are stable."},
            {"text": "Ship review package"},
        ],
    )

    document = Document(result["path"])
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Production Update",
        "Revenue and delivery are stable.",
        "Ship review package",
    ]
    assert document.paragraphs[0].style.name in {"Heading 1", "Title"}
    assert document.paragraphs[1].runs[0].bold is True
    assert document.paragraphs[2].style.name == "List Bullet"
    assert result["template_name"] == "Weekly Template"


def test_coordinator_learns_docx_template_then_generates_from_it(tmp_path):
    source = tmp_path / "weekly.docx"
    write_doc_template(source)
    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    document_store = DocumentTemplateStore(tmp_path / "doc_templates.json")
    llm = FakeLLM(
        json.dumps(
            {
                "format": "docx",
                "filename": "weekly-output.docx",
                "paragraphs": [
                    {"text": "Production Week 29"},
                    {"text": "No delivery blockers."},
                    {"text": "Send client package"},
                ],
            }
        )
    )
    coordinator = ArtifactCoordinator(
        generator,
        llm,
        document_template_store=document_store,
    )

    learned = coordinator.process(
        "learn this document format as template Weekly Template",
        attachments=[{"name": "weekly.docx", "sha256": "sha"}],
        file_paths=[source],
    )

    assert learned.matched is True
    assert learned.error_code is None
    assert learned.requested_format == "docx"
    assert "Weekly Template" in learned.message
    assert llm.calls == []

    generated = coordinator.process(
        "generate Word using template Weekly Template for production update"
    )

    assert generated.error_code is None
    assert generated.artifact["format"] == "docx"
    assert generated.artifact["name"] == "weekly-output.docx"

    document = Document(generated.artifact["path"])
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Production Week 29",
        "No delivery blockers.",
        "Send client package",
    ]
    assert document.paragraphs[0].style.name in {"Heading 1", "Title"}
    assert document.paragraphs[2].style.name == "List Bullet"
    assert "已保存的文档模板" in llm.calls[0]["system_prompt"]
