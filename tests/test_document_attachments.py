from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from docx import Document
from PIL import Image
import pytest


APP_ROOT = Path(__file__).resolve().parents[1] / "artpm_agent"

from artpm_agent.skills.skill_router import (
    MAX_DOCUMENT_FILE_SIZE,
    DocumentClassifierParser,
)


def test_document_parser_reads_real_pdf_docx_and_image(tmp_path):
    pdf_path = tmp_path / "brief.pdf"
    Image.new("RGB", (16, 12), "white").save(pdf_path, "PDF")

    docx_path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("项目说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "资产"
    table.cell(0, 1).text = "角色模型"
    document.save(docx_path)

    image_path = tmp_path / "reference.png"
    Image.new("RGB", (23, 17), "blue").save(image_path, "PNG")

    parser = DocumentClassifierParser()
    pdf_result = parser.run({"file_path": str(pdf_path)})
    docx_result = parser.run({"file_path": str(docx_path)})
    image_result = parser.run({"file_path": str(image_path)})

    assert pdf_result["success"] is True
    assert pdf_result["document_type"] == "PDF资料"
    assert pdf_result["extracted_data"]["page_count"] == 1

    assert docx_result["success"] is True
    assert docx_result["document_type"] == "Word资料"
    assert docx_result["extracted_data"] == {
        "paragraph_count": 1,
        "table_count": 1,
    }
    assert "项目说明" in docx_result["raw_text"]
    assert "资产\t角色模型" in docx_result["raw_text"]

    assert image_result["success"] is True
    assert image_result["document_type"] == "图片资料"
    assert image_result["requires_vision"] is True
    assert image_result["extracted_data"] == {
        "width": 23,
        "height": 17,
        "format": "PNG",
    }


def test_document_parser_uses_optional_unlimited_ocr_and_keeps_image_fallback(tmp_path):
    image_path = tmp_path / "ocr.png"
    Image.new("RGB", (23, 17), "white").save(image_path, "PNG")

    class SuccessfulOCR:
        def parse(self, images, **kwargs):
            assert images == [image_path.resolve()]
            assert kwargs["image_mode"] == "gundam"
            return SimpleNamespace(
                ok=True,
                text="# 交付清单\n角色原画 12 张",
                source="unlimited-ocr",
                degraded=False,
                truncated=False,
                error=None,
            )

    parsed = DocumentClassifierParser(
        {"unlimited_ocr_client": SuccessfulOCR()}
    ).run({"file_path": str(image_path)})

    assert parsed["success"] is True
    assert parsed["raw_text"].startswith("# 交付清单")
    assert parsed["requires_vision"] is False
    assert parsed["ocr_status"] == "completed"
    assert parsed["extracted_data"]["ocr"]["ok"] is True

    class FailedOCR:
        def parse(self, images, **kwargs):
            return SimpleNamespace(
                ok=False,
                text="",
                source="unlimited-ocr",
                degraded=True,
                truncated=False,
                error="service unavailable",
            )

    fallback = DocumentClassifierParser(
        {"unlimited_ocr_client": FailedOCR()}
    ).run({"file_path": str(image_path)})
    assert fallback["success"] is True
    assert fallback["raw_text"] == ""
    assert fallback["requires_vision"] is True
    assert fallback["ocr_status"] == "unavailable"
    assert fallback["extracted_data"]["ocr"]["error"] == "service unavailable"


def test_scanned_pdf_uses_multi_page_ocr_fallback(tmp_path):
    pdf_path = tmp_path / "scanned.pdf"
    Image.new("RGB", (32, 24), "white").save(pdf_path, "PDF")

    class SuccessfulPDFOCR:
        def parse(self, images, **kwargs):
            assert len(images) == 1
            assert Path(images[0]).is_file()
            assert kwargs["image_mode"] == "gundam"
            assert kwargs["prompt"] == "document parsing."
            return SimpleNamespace(
                ok=True,
                text="扫描合同正文",
                source="unlimited-ocr",
                degraded=False,
                truncated=False,
                error=None,
            )

    result = DocumentClassifierParser(
        {"unlimited_ocr_client": SuccessfulPDFOCR()}
    ).run({"file_path": str(pdf_path)})

    assert result["success"] is True
    assert "扫描合同正文" in result["raw_text"]
    assert result["ocr_status"] == "completed"
    assert result["extracted_data"]["ocr"]["rendered_pages"] == 1


def test_document_parser_validates_path_extension_empty_and_size(tmp_path):
    parser = DocumentClassifierParser()
    directory = tmp_path / "folder"
    directory.mkdir()
    unsupported = tmp_path / "payload.exe"
    unsupported.write_bytes(b"binary")
    empty = tmp_path / "empty.pdf"
    empty.touch()
    oversized = tmp_path / "oversized.pdf"
    with oversized.open("wb") as stream:
        stream.seek(MAX_DOCUMENT_FILE_SIZE)
        stream.write(b"x")

    missing_result = parser.run({"file_path": str(tmp_path / "missing.pdf")})
    directory_result = parser.run({"file_path": str(directory)})
    unsupported_result = parser.run({"file_path": str(unsupported)})
    empty_result = parser.run({"file_path": str(empty)})
    oversized_result = parser.run({"file_path": str(oversized)})

    assert missing_result["success"] is False
    assert "文件不存在" in missing_result["error"]
    assert directory_result["success"] is False
    assert "文件不存在" in directory_result["error"]
    assert unsupported_result["success"] is False
    assert "暂不支持" in unsupported_result["error"]
    assert empty_result["success"] is False
    assert "不能为空" in empty_result["error"]
    assert oversized_result["success"] is False
    assert "50 MB" in oversized_result["error"]


@pytest.mark.parametrize(
    ("filename", "contents"),
    [
        ("broken.pdf", b"not a pdf"),
        ("broken.docx", b"not a docx"),
        ("broken.png", b"not an image"),
    ],
)
def test_document_parser_rejects_corrupt_supported_formats(
    tmp_path,
    filename,
    contents,
):
    path = tmp_path / filename
    path.write_bytes(contents)

    result = DocumentClassifierParser().run({"file_path": str(path)})

    assert result["success"] is False
    assert result["error"]


def test_document_parser_rejects_image_content_extension_mismatch(tmp_path):
    disguised = tmp_path / "disguised.png"
    Image.new("RGB", (8, 8), "red").save(disguised, "JPEG")

    result = DocumentClassifierParser().run({"file_path": str(disguised)})

    assert result["success"] is False
    assert "与扩展名" in result["error"]
