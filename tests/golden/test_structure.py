"""结构性 golden 断言（ST-1..ST-7）。

本文件只断言**结构**，不含金额真值。真值断言见 `tests/golden/README.md` §6，
依赖业务方提供的 3 份历史项目（PRD §12.5，当前硬阻塞）。

未实现的能力在这里显式 skip 并打印原因，不得报告为已通过。
"""

from __future__ import annotations

import io
from hashlib import sha256
from pathlib import Path

import openpyxl
import pytest
from openpyxl import load_workbook

from artpm_agent.artifacts import WorkspaceArtifactGenerator
from artpm_agent.artifacts.xlsx_monthly import transform_monthly_workbook
from artpm_agent.parsers.excel_parser import ExcelQuoteParser
from artpm_agent.skills.cost_control_skill import CostControlSkill

COST_SOURCES = {"contract", "quote_history", "manual", "default"}


def _stream(rows, name="golden.xlsx") -> io.BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    stream.name = name
    workbook.save(stream)
    stream.seek(0)
    return stream


# ST-1：报价明细列名。缺 T1 报价单生成器，暂以解析器输出的资产字段为准。


def test_st1_asset_rows_match_input_count():
    """输入 N 个资产 → 解析出恰好 N 行明细。"""
    result = ExcelQuoteParser().parse(
        _stream(
            [
                ["资产名称", "数量", "单价", "总价"],
                ["主角", 2, 1000, 2000],
                ["配角", 1, 800, 800],
                ["场景", 3, 1500, 4500],
                ["道具", 1, 300, 300],
            ]
        )
    )
    assert result["document_type"] == "报价单"
    assert len(result["assets"]) == 4


@pytest.mark.skip(
    reason="T1 报价单生成器尚未实现（交付物模板库 §4），无法断言明细表列名"
)
def test_st1_quote_template_columns():
    raise AssertionError("unreachable")


# ST-2：费用构成算术自洽。


class _RateMember:
    def __init__(self, name, daily_cost, skill_level, cost_source="contract"):
        self.name = name
        self.daily_cost = daily_cost
        self.skill_level = skill_level
        self.cost_source = cost_source
        self.cost_effective_at = None


class _RateDB:
    def __init__(self, members):
        self._members = list(members)

    def list_members(self, is_active=None):
        return list(self._members)


def _rate_skill(overhead_rate=0.15, tax_rate=0.06):
    return CostControlSkill(
        {
            "config": {
                "cost_config": {
                    "overhead_rate": overhead_rate,
                    "tax_rate": tax_rate,
                }
            },
            "database": _RateDB([_RateMember("张三", 500, "中级")]),
        }
    )


def test_st2_fee_composition_is_arithmetic_consistent():
    result = _rate_skill().estimate(24, "中级", 2)

    assert result["success"] is True
    labor = result["labor_cost"]
    overhead = result["overhead_cost"]
    tax = result["tax_cost"]
    total = result["total_cost"]

    assert labor + overhead + tax == pytest.approx(total, abs=0.01)
    assert labor * (1 + result["overhead_rate"]) * (1 + result["tax_rate"]) == (
        pytest.approx(total, abs=0.01)
    )


def test_st2_unknown_staff_level_is_rejected_not_guessed():
    result = _rate_skill().estimate(8, "不存在档位", 1)
    assert result["success"] is False


# ST-3：费率来源必须可追溯。列已由 0005 迁移落到 team_members，
# 因此这条不再需要跳过。


def test_st3_default_cost_must_be_labelled():
    """从配置播种的费率必须带 cost_source='default'，不能只有数字。"""
    from artpm_agent.skills.rate_model import (
        VALID_COST_SOURCES,
        resolve_member_rate,
    )

    seeded = {
        "name": "张三",
        "daily_cost": 500,
        "cost_source": "default",
        "cost_effective_at": "2026-09-19",
    }
    result = resolve_member_rate(seeded)

    assert result.resolved is True
    assert result.cost_source == "default"
    assert result.cost_source in VALID_COST_SOURCES
    assert result.cost_effective_at is not None


def test_st3_rate_without_source_is_reported_as_missing():
    """有数字没来源 = 不可追溯，必须标记而不是当正常费率用。"""
    from artpm_agent.skills.rate_model import resolve_member_rate

    result = resolve_member_rate({"name": "李四", "daily_cost": 500})

    assert result.daily_cost == 500.0
    assert result.resolved is False
    assert [item.field for item in result.missing] == ["cost_source"]


def test_st3_rate_model_reads_orm_cost_source_column():
    """确认列真的在 ORM 映射上，而不只是迁移写了 DDL。"""
    from artpm_agent.database.models import TeamMember

    assert "cost_source" in TeamMember.__table__.columns
    assert "cost_effective_at" in TeamMember.__table__.columns
    assert "capacity_days_per_month" in TeamMember.__table__.columns
    assert "contact_wecom" in TeamMember.__table__.columns
    # 不允许模型级默认值伪造来源
    assert TeamMember.__table__.columns["cost_source"].default is None


# ST-4：缺费率返回 needs_input，不产出报价。


def test_st4_missing_daily_cost_returns_needs_input():
    """技能层：人员没有 daily_cost 时给缺项清单，不猜值、不出价。"""
    skill = CostControlSkill(
        {
            "config": {"cost_config": {"overhead_rate": 0.15, "tax_rate": 0.06}},
            "database": _RateDB([_RateMember("王五", None, "中级")]),
        }
    )
    result = skill.estimate(8, "中级", 1)

    assert result["success"] is False
    assert result["status"] == "needs_input"
    assert result["missing"][0]["field"] == "daily_cost"
    assert result["subjects"] == ["王五"]
    assert "total_cost" not in result


def test_st4_no_rate_table_is_needs_input_not_builtin_default():
    """没有费率表连估算都不能给，内置常量必须已经删干净。"""
    from artpm_agent.skills import cost_control_skill as module

    assert not hasattr(module, "DEFAULT_STAFF_LEVELS")
    assert not hasattr(module, "DEFAULT_OVERHEAD_RATE")
    assert not hasattr(module, "DEFAULT_TAX_RATE")

    result = CostControlSkill(
        {"config": {"cost_config": {"overhead_rate": 0.15, "tax_rate": 0.06}}}
    ).estimate(8, "高级", 1)

    assert result["status"] == "needs_input"
    assert result["missing"][0]["field"] == "team_members"


def test_st4_rate_model_returns_needs_input_without_rate():
    """领域层已具备该行为：缺费率时给出缺项清单，不猜值。"""
    from artpm_agent.skills.rate_model import needs_input, resolve_member_rate

    resolution = resolve_member_rate({"name": "王五"})
    payload = needs_input(resolution.missing)

    assert payload["status"] == "needs_input"
    assert payload["success"] is False
    assert payload["subjects"] == ["王五"]
    assert payload["missing"][0]["field"] == "daily_cost"


# ST-5：生成的 .xlsx 可重新打开，行数与请求一致，SHA-256 与发布副本匹配。


def test_st5_generated_monthly_workbook_reopens_and_matches_sha(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "人天分配"
    sheet.append(["任务", "月份", "商务人天"])
    sheet.append(["麦迪", "2026-11", 18])
    sheet.append(["裸模", "2026-12", 15])
    workbook.save(source)
    workbook.close()

    generator = WorkspaceArtifactGenerator(tmp_path / "artifacts")
    result = generator.generate_xlsx_monthly("11月汇总.xlsx", source, "2026-11")

    path = Path(result["path"])
    assert path.is_file()
    assert result["verification"]["status"] == "passed"
    assert result["sha256"] == sha256(path.read_bytes()).hexdigest()

    reopened = load_workbook(path, read_only=True)
    assert "人天分配" in reopened.sheetnames
    assert "处理说明" in reopened.sheetnames
    assert reopened["人天分配"].max_row == 2
    reopened.close()


# ST-6：同一输入两次运行结构一致。


def test_st6_monthly_transform_is_deterministic(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "人天分配"
    sheet.append(["任务", "月份", "商务人天"])
    sheet.append(["麦迪", "2026-11", 18])
    sheet.append(["裸模", "2026-11", 15])
    sheet.append(["场景", "2026-12", 12])
    workbook.save(source)
    workbook.close()

    first = transform_monthly_workbook(source, tmp_path / "a.xlsx", "2026-11")
    second = transform_monthly_workbook(source, tmp_path / "b.xlsx", "2026-11")

    assert first.sheets == second.sheets
    assert first.matched_rows == second.matched_rows
    assert first.month_column == second.month_column

    def read_rows(path):
        book = load_workbook(path)
        rows = [tuple(row) for row in book["人天分配"].iter_rows(values_only=True)]
        book.close()
        return rows

    assert read_rows(tmp_path / "a.xlsx") == read_rows(tmp_path / "b.xlsx")


# ST-7：Excel 分类不再一律返回「报价单」。


@pytest.mark.parametrize(
    "rows,forbidden",
    [
        ([["任务", "商务人天", "已分配人天"], ["麦迪", 18, 0]], "报价单"),
        ([["人员", "月度产能"], ["张三", 20]], "报价单"),
    ],
)
def test_st7_non_quote_sheets_are_not_labelled_as_quotes(rows, forbidden):
    result = ExcelQuoteParser().parse(_stream(rows))
    assert result["document_type"] != forbidden
    assert result["document_type"] in {
        "人天分配表",
        "产能表",
        "排期表",
        "unknown",
    }
