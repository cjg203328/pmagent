"""In-memory editable document model for one-sentence editing.

Holds a parsed spreadsheet or Word document so it can be mutated in memory,
previewed, and serialized back to bytes. The original file is NEVER mutated -
callers persist changes through the versioned artifact generator instead.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook

_EXCEL_EXT = {".xlsx", ".xls"}
_WORD_EXT = {".docx"}


def _safe_sheet_name(name: str, limit: int = 31) -> str:
    invalid = set('[]:*?/\\')
    cleaned = "".join("_" if ch in invalid else ch for ch in (name or "Sheet1"))
    cleaned = cleaned[:limit].strip() or "Sheet1"
    return cleaned


class EditableDocument:
    """A spreadsheet or Word document loaded for in-memory editing."""

    def __init__(self, path: str | Path, kind: str) -> None:
        self.path = Path(path)
        self.kind = kind  # "excel" | "word"
        self.original_name = self.path.name
        # excel
        self.sheets: dict[str, "pd.DataFrame"] = {}
        self.active_sheet = ""
        # word
        self.blocks: list[dict[str, Any]] = []
        self._original_bytes = self.path.read_bytes()

    # ---- construction -------------------------------------------------
    @classmethod
    def load(cls, path: str | Path) -> "EditableDocument":
        p = Path(path)
        ext = p.suffix.lower()
        if ext in _EXCEL_EXT:
            doc = cls(p, "excel")
            doc._load_excel()
        elif ext == ".docx":
            doc = cls(p, "word")
            doc._load_word()
        else:
            raise ValueError(f"unsupported edit format: {ext}")
        return doc

    def _load_excel(self) -> None:
        with pd.ExcelFile(self.path) as xls:
            for name in xls.sheet_names:
                df = xls.parse(name, dtype=object)
                df = df.where(pd.notnull(df), None)
                df.columns = [str(c) for c in df.columns]
                self.sheets[name] = df
            self.active_sheet = xls.sheet_names[0] if xls.sheet_names else "Sheet1"

    def _load_word(self) -> None:
        document = Document(self.path)
        blocks: list[dict[str, Any]] = []
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, document)
                text = para.text
                style_name = (para.style.name if para.style else "") or ""
                if style_name.startswith("Heading"):
                    level = 1
                    digits = "".join(ch for ch in style_name if ch.isdigit())
                    if digits:
                        level = int(digits)
                    blocks.append({"type": "heading", "text": text, "level": level})
                else:
                    blocks.append({"type": "paragraph", "text": text})
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                rows: list[list[Any]] = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                blocks.append({"type": "table", "rows": rows})
        self.blocks = blocks

    # ---- accessors ----------------------------------------------------
    @property
    def active_df(self) -> "pd.DataFrame":
        return self.sheets.get(self.active_sheet, pd.DataFrame())

    def sheet_names(self) -> list[str]:
        return list(self.sheets.keys())

    def summary(self) -> str:
        if self.kind == "excel":
            df = self.active_df
            return (
                f"{self.original_name} - sheet '{self.active_sheet}'"
                f" - {df.shape[0]} rows x {df.shape[1]} cols"
            )
        return f"{self.original_name} - Word - {len(self.blocks)} blocks"

    # ---- (de)serialization -------------------------------------------
    def to_bytes(self) -> bytes:
        if self.kind == "excel":
            return self._excel_bytes()
        return self._word_bytes()

    def _excel_bytes(self) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)
        for name, df in self.sheets.items():
            ws = wb.create_sheet(title=_safe_sheet_name(name))
            ws.append(list(df.columns))
            for _, row in df.iterrows():
                ws.append([None if v is None else v for v in row.tolist()])
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()

    def _word_bytes(self) -> bytes:
        document = Document()
        for block in self.blocks:
            btype = block.get("type")
            if btype == "heading":
                document.add_heading(block.get("text", ""), level=block.get("level", 1))
            elif btype == "paragraph":
                document.add_paragraph(block.get("text", ""))
            elif btype == "table":
                rows = block.get("rows", [])
                if not rows:
                    continue
                cols = max((len(r) for r in rows), default=0)
                table = document.add_table(rows=len(rows), cols=cols)
                for i, row in enumerate(rows):
                    for j, value in enumerate(row):
                        cell = table.cell(i, j)
                        cell.text = "" if value is None else str(value)
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def diff_from(self, other: "EditableDocument") -> list[str]:
        """Human-readable list of changes vs *other* (the prior state)."""
        changes: list[str] = []
        if self.kind == "excel":
            before = other.active_df
            after = self.active_df
            cols = list(after.columns)
            n = min(len(before), len(after))
            for i in range(n):
                for j, col in enumerate(cols):
                    bv = before.iloc[i, j] if j < before.shape[1] else None
                    av = after.iloc[i, j] if j < after.shape[1] else None
                    if str(bv) != str(av):
                        changes.append(
                            f"row {i + 1} | {col}: {_short(bv)} -> {_short(av)}"
                        )
            if len(after) > len(before):
                changes.append(f"+{len(after) - len(before)} rows")
            elif len(after) < len(before):
                changes.append(f"-{len(before) - len(after)} rows")
        else:
            for i, block in enumerate(self.blocks):
                ob = other.blocks[i] if i < len(other.blocks) else {}
                if block.get("text") != ob.get("text"):
                    changes.append(
                        f"block {i + 1} ({block.get('type')}): "
                        f"{_short(ob.get('text'))} -> {_short(block.get('text'))}"
                    )
        return changes


def _short(value: Any, limit: int = 40) -> str:
    if value is None:
        return "(empty)"
    text = str(value)
    return text if len(text) <= limit else text[:limit] + "..."
