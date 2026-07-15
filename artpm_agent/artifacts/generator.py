"""Generate new XLSX and DOCX files below one workspace artifact root."""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from io import BytesIO, StringIO
import csv
from datetime import date, datetime, timezone
from hashlib import sha256
import math
from pathlib import Path
import re
import shutil
from typing import Any
import unicodedata
from uuid import uuid4

from docx import Document
from openpyxl import Workbook
from openpyxl import load_workbook

from artpm_agent.utils.multimodal_markdown import LocalMarkdownConverter


_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_WINDOWS_DRIVE = re.compile(r"^[A-Za-z]:")
_WINDOWS_RESERVED_NAMES = frozenset(
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{index}" for index in range(1, 10)}
    | {f"LPT{index}" for index in range(1, 10)}
)
_INVALID_SHEET_NAME = re.compile(r"[\\/*?:\[\]]")
_ALLOWED_CELL_TYPES = (str, int, float, bool, date, datetime)


class WorkspaceArtifactGenerator:
    """Create versioned artifacts without exposing overwrite or delete operations."""

    MAX_XLSX_ROWS = 1_048_575
    MAX_XLSX_COLUMNS = 16_384

    def __init__(
        self,
        artifacts_root: str | Path,
        *,
        max_rows: int = 10_000,
        max_columns: int = 100,
        max_paragraphs: int = 2_000,
        max_cell_chars: int = 32_000,
        max_total_text_chars: int = 1_000_000,
        max_file_size: int = 50 * 1024 * 1024,
        max_versions: int = 10_000,
    ) -> None:
        root = Path(artifacts_root).expanduser().resolve()
        root.mkdir(parents=True, exist_ok=True)
        if not root.is_dir():
            raise ValueError(f"Artifact root is not a directory: {root}")
        self.root = root
        self.max_rows = self._bounded_positive(
            max_rows,
            "max_rows",
            self.MAX_XLSX_ROWS,
        )
        self.max_columns = self._bounded_positive(
            max_columns,
            "max_columns",
            self.MAX_XLSX_COLUMNS,
        )
        self.max_paragraphs = self._positive(max_paragraphs, "max_paragraphs")
        self.max_cell_chars = self._positive(max_cell_chars, "max_cell_chars")
        self.max_total_text_chars = self._positive(
            max_total_text_chars,
            "max_total_text_chars",
        )
        self.max_file_size = self._positive(max_file_size, "max_file_size")
        self.max_versions = self._positive(max_versions, "max_versions")

    @staticmethod
    def _positive(value: int, field: str) -> int:
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            raise ValueError(f"{field} must be a positive integer")
        return value

    @classmethod
    def _bounded_positive(cls, value: int, field: str, maximum: int) -> int:
        value = cls._positive(value, field)
        if value > maximum:
            raise ValueError(f"{field} cannot exceed {maximum}")
        return value

    @staticmethod
    def _sequence(value: Any, field: str) -> Sequence[Any]:
        if isinstance(value, (str, bytes, bytearray)) or not isinstance(
            value, Sequence
        ):
            raise ValueError(f"{field} must be a sequence")
        return value

    @staticmethod
    def _safe_filename(filename: str, extension: str) -> str:
        if not isinstance(filename, str) or not filename.strip():
            raise ValueError("filename must be a non-empty string")
        normalized = unicodedata.normalize("NFKC", filename.strip())
        if (
            "/" in normalized
            or "\\" in normalized
            or normalized in {".", ".."}
            or _WINDOWS_DRIVE.match(normalized)
        ):
            raise ValueError("filename must not contain a path")

        suffix = Path(normalized).suffix
        expected_suffix = f".{extension}"
        if suffix and suffix.casefold() != expected_suffix:
            raise ValueError(f"filename extension must be {expected_suffix}")
        stem = normalized[: -len(suffix)] if suffix else normalized
        stem = _INVALID_FILENAME.sub("_", stem).strip(" .")
        if not stem:
            raise ValueError("filename has no usable characters")
        if stem.upper() in _WINDOWS_RESERVED_NAMES:
            stem = f"_{stem}"

        max_stem_length = 120 - len(expected_suffix)
        stem = stem[:max_stem_length].rstrip(" .")
        if not stem:
            raise ValueError("filename has no usable characters")
        return f"{stem}{expected_suffix}"

    @staticmethod
    def _sheet_name(value: Any) -> str:
        if value is None:
            return "Sheet1"
        if not isinstance(value, str) or not value.strip():
            raise ValueError("sheet_name must be a non-empty string")
        name = _INVALID_SHEET_NAME.sub("_", value.strip())[:31].strip("'")
        if not name:
            raise ValueError("sheet_name has no usable characters")
        return name

    def _safe_path(self, filename: str) -> Path:
        candidate = (self.root / filename).resolve()
        try:
            candidate.relative_to(self.root)
        except ValueError as error:
            raise ValueError("Artifact path escapes its workspace root") from error
        if candidate.parent != self.root:
            raise ValueError("Artifacts must be direct children of the workspace root")
        return candidate

    def artifact_path(self, stored_path: str) -> Path:
        """Resolve one direct-child artifact path below this workspace root."""
        if not isinstance(stored_path, str) or not stored_path.strip():
            raise ValueError("stored_path must be a non-empty string")
        if "/" in stored_path or "\\" in stored_path or _WINDOWS_DRIVE.match(stored_path):
            raise ValueError("stored_path must be a direct artifact filename")
        path = (self.root / stored_path).resolve()
        try:
            path.relative_to(self.root)
        except ValueError as error:
            raise ValueError("Artifact path escapes its workspace root") from error
        if path.parent != self.root or not path.is_file():
            raise ValueError("artifact does not exist")
        if path.stat().st_size > self.max_file_size:
            raise ValueError("artifact exceeds the configured file size limit")
        return path

    @staticmethod
    def _versioned_name(filename: str, version: int) -> str:
        path = Path(filename)
        if version == 1:
            return filename
        return f"{path.stem} ({version}){path.suffix}"

    def _publish_new(self, temporary_path: Path, filename: str) -> tuple[Path, int]:
        for version in range(1, self.max_versions + 1):
            destination = self._safe_path(self._versioned_name(filename, version))
            if destination.exists():
                continue
            try:
                with temporary_path.open("rb") as source, destination.open("xb") as target:
                    shutil.copyfileobj(source, target, length=1024 * 1024)
            except FileExistsError:
                continue
            except IsADirectoryError:
                continue
            except Exception:
                destination.unlink(missing_ok=True)
                raise
            return destination, version
        raise FileExistsError(
            f"No free artifact version found for {filename!r} "
            f"after {self.max_versions} attempts"
        )

    def _finalize(
        self,
        temporary_path: Path,
        filename: str,
        artifact_format: str,
        mime_type: str,
        details: Mapping[str, Any],
    ) -> dict[str, Any]:
        size = temporary_path.stat().st_size
        if size <= 0:
            raise RuntimeError("Generated artifact is empty")
        if size > self.max_file_size:
            raise ValueError(
                f"Generated artifact exceeds {self.max_file_size} bytes"
            )
        digest = sha256(temporary_path.read_bytes()).hexdigest()
        destination, version = self._publish_new(temporary_path, filename)
        preview = self.preview_path(destination)
        return {
            "id": uuid4().hex,
            "name": destination.name,
            "stored_path": destination.name,
            "path": str(destination),
            "format": artifact_format,
            "mime_type": mime_type,
            "size": size,
            "sha256": digest,
            "version": version,
            "created_at": datetime.now(timezone.utc).isoformat(
                timespec="microseconds"
            ),
            "preview_markdown": preview.get("preview_markdown", ""),
            "export_formats": self.available_export_formats(artifact_format),
            **dict(details),
        }

    def _temporary_path(self) -> Path:
        return self._safe_path(f".artifact-{uuid4().hex}.tmp")

    @staticmethod
    def available_export_formats(artifact_format: str) -> list[str]:
        """Return browser-safe export formats for a generated artifact."""
        normalized = str(artifact_format or "").lower().lstrip(".")
        if normalized == "xlsx":
            return ["csv", "md", "txt", "docx"]
        if normalized == "docx":
            return ["md", "txt"]
        if normalized == "csv":
            return ["md", "txt", "xlsx"]
        if normalized in {"md", "txt"}:
            return ["docx"]
        return []

    @staticmethod
    def _mime_type_for_format(target_format: str) -> str:
        return {
            "csv": "text/csv",
            "docx": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            "md": "text/markdown",
            "txt": "text/plain",
            "xlsx": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        }.get(target_format, "application/octet-stream")

    def preview_path(self, path: str | Path, *, max_chars: int = 4_000) -> dict[str, Any]:
        """Create a bounded Markdown preview from a trusted local artifact path."""
        artifact_path = Path(path).expanduser().resolve()
        try:
            artifact_path.relative_to(self.root)
        except ValueError as error:
            raise ValueError("Artifact path escapes its workspace root") from error
        if artifact_path.parent != self.root or not artifact_path.is_file():
            raise ValueError("artifact does not exist")
        converter = LocalMarkdownConverter(max_chars=max(max_chars, 1024))
        converted = converter.convert(artifact_path)
        preview_markdown = converted.markdown[:max_chars].strip()
        return {
            "success": converted.success,
            "name": artifact_path.name,
            "format": artifact_path.suffix.lower().lstrip("."),
            "preview_markdown": preview_markdown,
            "truncated": converted.truncated or len(converted.markdown) > max_chars,
            "error": converted.error,
        }

    def preview_artifact(
        self,
        stored_path: str,
        *,
        max_chars: int = 4_000,
    ) -> dict[str, Any]:
        """Create a bounded Markdown preview for a stored generated artifact."""
        return self.preview_path(self.artifact_path(stored_path), max_chars=max_chars)

    @staticmethod
    def _markdown_to_docx_bytes(markdown: str) -> bytes:
        document = Document()
        lines = str(markdown or "").splitlines()
        has_content = False
        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                continue
            has_content = True
            if line.startswith("#"):
                hashes = len(line) - len(line.lstrip("#"))
                text = line[hashes:].strip()
                document.add_heading(text or "Section", level=max(1, min(hashes, 9)))
            elif line.startswith("- "):
                document.add_paragraph(line[2:].strip(), style="List Bullet")
            elif line[:3].replace(".", "").isdigit() and ". " in line[:5]:
                document.add_paragraph(line.split(". ", 1)[1], style="List Number")
            else:
                document.add_paragraph(line)
        if not has_content:
            document.add_paragraph("")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    @staticmethod
    def _xlsx_to_csv_bytes(path: Path) -> bytes:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            stream = StringIO(newline="")
            writer = csv.writer(stream)
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
            return stream.getvalue().encode("utf-8-sig")
        finally:
            workbook.close()

    @staticmethod
    def _csv_to_xlsx_bytes(path: Path) -> bytes:
        raw = path.read_bytes()
        text = ""
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = raw.decode("utf-8", errors="replace")
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(title="Sheet1")
        for row in csv.reader(text.splitlines()):
            worksheet.append(row)
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def export_artifact_bytes(
        self,
        stored_path: str,
        target_format: str,
    ) -> dict[str, Any]:
        """Return converted bytes for browser download without mutating storage."""
        source = self.artifact_path(stored_path)
        source_format = source.suffix.lower().lstrip(".")
        target = str(target_format or "").lower().lstrip(".")
        if target == source_format:
            data = source.read_bytes()
        elif target == "csv" and source_format == "xlsx":
            data = self._xlsx_to_csv_bytes(source)
        elif target in {"md", "txt"}:
            converted = LocalMarkdownConverter(max_chars=1_000_000).convert(source)
            if not converted.success:
                raise ValueError(converted.error or "artifact cannot be exported")
            data = converted.markdown.encode("utf-8")
        elif target == "docx" and source_format in {"xlsx", "csv", "md", "txt"}:
            if source_format in {"md", "txt"}:
                markdown = source.read_text(encoding="utf-8", errors="replace")
            else:
                converted = LocalMarkdownConverter(max_chars=1_000_000).convert(source)
                if not converted.success:
                    raise ValueError(converted.error or "artifact cannot be exported")
                markdown = converted.markdown
            data = self._markdown_to_docx_bytes(markdown)
        elif target == "xlsx" and source_format == "csv":
            data = self._csv_to_xlsx_bytes(source)
        else:
            raise ValueError(
                f"cannot export {source_format or 'artifact'} as {target or 'unknown'}"
            )
        if len(data) > self.max_file_size:
            raise ValueError("exported artifact exceeds the configured file size limit")
        filename = f"{source.stem}.{target}"
        return {
            "filename": filename,
            "format": target,
            "mime_type": self._mime_type_for_format(target),
            "data": data,
            "size": len(data),
        }

    def export_artifact(
        self,
        stored_path: str,
        target_format: str,
    ) -> dict[str, Any]:
        """Save a converted copy as a new versioned artifact."""
        exported = self.export_artifact_bytes(stored_path, target_format)
        temporary_path = self._temporary_path()
        try:
            temporary_path.write_bytes(exported["data"])
            return self._finalize(
                temporary_path,
                exported["filename"],
                exported["format"],
                exported["mime_type"],
                {"source_artifact": stored_path},
            )
        finally:
            temporary_path.unlink(missing_ok=True)

    def _cell_value(self, value: Any, field: str) -> Any:
        if value is None:
            return None
        if not isinstance(value, _ALLOWED_CELL_TYPES):
            raise ValueError(f"{field} contains an unsupported cell value")
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"{field} contains a non-finite number")
        if isinstance(value, str):
            if len(value) > self.max_cell_chars:
                raise ValueError(
                    f"{field} exceeds the {self.max_cell_chars} character limit"
                )
            if value.startswith("="):
                return f"'{value}"
        return value

    def _table_rows(
        self,
        columns: tuple[str, ...],
        rows: Iterable[Any],
    ) -> Iterable[list[Any]]:
        if isinstance(rows, (str, bytes, bytearray, Mapping)) or not isinstance(
            rows, Iterable
        ):
            raise ValueError("table.rows must be an iterable of rows")
        column_set = set(columns)
        for index, row in enumerate(rows, start=1):
            if index > self.max_rows:
                raise ValueError(f"table.rows cannot exceed {self.max_rows} rows")
            if isinstance(row, Mapping):
                unknown = set(row) - column_set
                if unknown:
                    raise ValueError(
                        "table row contains unknown columns: "
                        + ", ".join(sorted(str(item) for item in unknown))
                    )
                values = [row.get(column) for column in columns]
            else:
                values = list(self._sequence(row, f"table.rows[{index - 1}]"))
                if len(values) != len(columns):
                    raise ValueError(
                        f"table.rows[{index - 1}] must contain {len(columns)} cells"
                    )
            yield [
                self._cell_value(value, f"table.rows[{index - 1}]")
                for value in values
            ]

    def generate_xlsx(
        self,
        filename: str,
        table: Mapping[str, Any],
    ) -> dict[str, Any]:
        """Generate one new workbook from columns and row mappings/sequences."""
        if not isinstance(table, Mapping):
            raise ValueError("table must be a mapping")
        unknown = set(table) - {"columns", "rows", "sheet_name"}
        if unknown:
            raise ValueError(
                "table contains unsupported fields: "
                + ", ".join(sorted(str(item) for item in unknown))
            )
        raw_columns = self._sequence(table.get("columns"), "table.columns")
        if not raw_columns or len(raw_columns) > self.max_columns:
            raise ValueError(
                f"table.columns must contain 1 to {self.max_columns} columns"
            )
        columns: list[str] = []
        header_values: list[str] = []
        for index, column in enumerate(raw_columns):
            if not isinstance(column, str) or not column.strip():
                raise ValueError(f"table.columns[{index}] must be non-empty text")
            column = column.strip()
            columns.append(column)
            header_values.append(
                self._cell_value(column, f"table.columns[{index}]")
            )
        if len(columns) != len(set(columns)):
            raise ValueError("table.columns must be unique")

        safe_name = self._safe_filename(filename, "xlsx")
        sheet_name = self._sheet_name(table.get("sheet_name"))
        normalized_rows = list(
            self._table_rows(tuple(columns), table.get("rows", []))
        )
        temporary_path = self._temporary_path()
        try:
            workbook = Workbook(write_only=True)
            worksheet = workbook.create_sheet(title=sheet_name)
            worksheet.append(header_values)
            for values in normalized_rows:
                worksheet.append(values)
            workbook.save(temporary_path)
            return self._finalize(
                temporary_path,
                safe_name,
                "xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                {
                    "rows": len(normalized_rows),
                    "columns": len(columns),
                    "sheet_name": sheet_name,
                },
            )
        finally:
            temporary_path.unlink(missing_ok=True)

    def _normalized_columns(
        self,
        raw_columns: Sequence[Any],
        field: str,
    ) -> tuple[list[str], list[str]]:
        if not raw_columns or len(raw_columns) > self.max_columns:
            raise ValueError(
                f"{field} must contain 1 to {self.max_columns} columns"
            )
        columns: list[str] = []
        header_values: list[str] = []
        for index, column in enumerate(raw_columns):
            if not isinstance(column, str) or not column.strip():
                raise ValueError(f"{field}[{index}] must be non-empty text")
            column = column.strip()
            columns.append(column)
            header_values.append(self._cell_value(column, f"{field}[{index}]"))
        if len(columns) != len(set(columns)):
            raise ValueError(f"{field} must be unique")
        return columns, header_values

    @staticmethod
    def _template_rows_for_sheet(
        rows_by_sheet: Mapping[str, Any] | Sequence[Any] | None,
        sheet_name: str,
        sheet_index: int,
    ) -> Iterable[Any]:
        if rows_by_sheet is None:
            return []
        if isinstance(rows_by_sheet, Mapping):
            return (
                rows_by_sheet.get(sheet_name)
                or rows_by_sheet.get(str(sheet_index))
                or rows_by_sheet.get(sheet_index)
                or []
            )
        if isinstance(rows_by_sheet, Sequence) and not isinstance(
            rows_by_sheet,
            (str, bytes, bytearray),
        ):
            if sheet_index < len(rows_by_sheet):
                return rows_by_sheet[sheet_index] or []
            return []
        raise ValueError("rows_by_sheet must be a mapping, sequence, or None")

    def generate_xlsx_from_template(
        self,
        filename: str,
        template: Mapping[str, Any],
        *,
        rows_by_sheet: Mapping[str, Any] | Sequence[Any] | None = None,
    ) -> dict[str, Any]:
        """Generate a workbook that preserves saved worksheet names and columns."""
        if not isinstance(template, Mapping):
            raise ValueError("template must be a mapping")
        raw_sheets = self._sequence(template.get("sheets"), "template.sheets")
        if not raw_sheets:
            raise ValueError("template.sheets must contain at least one sheet")

        safe_name = self._safe_filename(filename, "xlsx")
        temporary_path = self._temporary_path()
        total_rows = 0
        max_columns = 0
        sheet_details: list[dict[str, Any]] = []
        try:
            workbook = Workbook(write_only=True)
            for sheet_index, raw_sheet in enumerate(raw_sheets):
                if not isinstance(raw_sheet, Mapping):
                    raise ValueError(
                        f"template.sheets[{sheet_index}] must be a mapping"
                    )
                sheet_name = self._sheet_name(
                    raw_sheet.get("name") or f"Sheet{sheet_index + 1}"
                )
                raw_columns = self._sequence(
                    raw_sheet.get("columns"),
                    f"template.sheets[{sheet_index}].columns",
                )
                columns, header_values = self._normalized_columns(
                    raw_columns,
                    f"template.sheets[{sheet_index}].columns",
                )
                normalized_rows = list(
                    self._table_rows(
                        tuple(columns),
                        self._template_rows_for_sheet(
                            rows_by_sheet,
                            sheet_name,
                            sheet_index,
                        ),
                    )
                )
                worksheet = workbook.create_sheet(title=sheet_name)
                worksheet.append(header_values)
                for values in normalized_rows:
                    worksheet.append(values)
                total_rows += len(normalized_rows)
                max_columns = max(max_columns, len(columns))
                sheet_details.append(
                    {
                        "name": sheet_name,
                        "rows": len(normalized_rows),
                        "columns": len(columns),
                    }
                )
            workbook.save(temporary_path)
            return self._finalize(
                temporary_path,
                safe_name,
                "xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                {
                    "rows": total_rows,
                    "columns": max_columns,
                    "sheet_name": sheet_details[0]["name"],
                    "sheets": sheet_details,
                    "template_id": template.get("id"),
                    "template_name": template.get("name"),
                },
            )
        finally:
            temporary_path.unlink(missing_ok=True)

    def _paragraph(self, value: Any, index: int) -> dict[str, Any]:
        if isinstance(value, str):
            paragraph: dict[str, Any] = {"text": value}
        elif isinstance(value, Mapping):
            paragraph = dict(value)
        else:
            raise ValueError(f"paragraphs[{index}] must be text or a mapping")
        unknown = set(paragraph) - {"text", "kind", "level", "bold", "italic"}
        if unknown:
            raise ValueError(
                f"paragraphs[{index}] contains unsupported fields: "
                + ", ".join(sorted(str(item) for item in unknown))
            )
        text = paragraph.get("text")
        if not isinstance(text, str):
            raise ValueError(f"paragraphs[{index}].text must be text")
        if len(text) > self.max_total_text_chars:
            raise ValueError(f"paragraphs[{index}].text is too large")
        kind = paragraph.get("kind", "paragraph")
        if kind not in {"paragraph", "heading", "bullet", "numbered"}:
            raise ValueError(f"paragraphs[{index}].kind is unsupported")
        level = paragraph.get("level", 1)
        if isinstance(level, bool) or not isinstance(level, int) or not 1 <= level <= 9:
            raise ValueError(f"paragraphs[{index}].level must be between 1 and 9")
        for key in ("bold", "italic"):
            if key in paragraph and not isinstance(paragraph[key], bool):
                raise ValueError(f"paragraphs[{index}].{key} must be boolean")
        return {
            "text": text,
            "kind": kind,
            "level": level,
            "bold": paragraph.get("bold", False),
            "italic": paragraph.get("italic", False),
        }

    def _template_paragraphs(
        self,
        template: Mapping[str, Any],
        paragraphs: Sequence[Any] | None,
    ) -> list[dict[str, Any]]:
        raw_outline = self._sequence(
            template.get("paragraphs"),
            "template.paragraphs",
        )
        if not raw_outline:
            raise ValueError("template.paragraphs must contain at least one item")
        outline = [
            self._paragraph(item, index)
            for index, item in enumerate(raw_outline[: self.max_paragraphs])
        ]

        if paragraphs is None:
            return outline

        raw_paragraphs = self._sequence(paragraphs, "paragraphs")
        if not raw_paragraphs:
            return outline

        normalized: list[dict[str, Any]] = []
        for index, raw in enumerate(raw_paragraphs):
            incoming = self._paragraph(raw, index)
            if index < len(outline):
                style = outline[index]
                incoming.update(
                    {
                        "kind": style["kind"],
                        "level": style["level"],
                        "bold": style["bold"] or incoming["bold"],
                        "italic": style["italic"] or incoming["italic"],
                    }
                )
            normalized.append(incoming)
        return normalized

    def generate_docx_from_template(
        self,
        filename: str,
        template: Mapping[str, Any],
        *,
        paragraphs: Sequence[Any] | None = None,
    ) -> dict[str, Any]:
        """Generate a document while preserving saved paragraph structure."""
        if not isinstance(template, Mapping):
            raise ValueError("template must be a mapping")
        normalized = self._template_paragraphs(template, paragraphs)
        artifact = self.generate_docx(filename, normalized)
        artifact["template_id"] = template.get("id")
        artifact["template_name"] = template.get("name")
        return artifact

    def generate_docx(
        self,
        filename: str,
        paragraphs: Sequence[Any],
    ) -> dict[str, Any]:
        """Generate one new document from structured paragraphs."""
        raw_paragraphs = self._sequence(paragraphs, "paragraphs")
        if not raw_paragraphs or len(raw_paragraphs) > self.max_paragraphs:
            raise ValueError(
                f"paragraphs must contain 1 to {self.max_paragraphs} items"
            )
        normalized = [
            self._paragraph(value, index)
            for index, value in enumerate(raw_paragraphs)
        ]
        total_chars = sum(len(item["text"]) for item in normalized)
        if total_chars > self.max_total_text_chars:
            raise ValueError(
                "paragraph text exceeds the total character limit of "
                f"{self.max_total_text_chars}"
            )

        safe_name = self._safe_filename(filename, "docx")
        temporary_path = self._temporary_path()
        try:
            document = Document()
            for item in normalized:
                if item["kind"] == "heading":
                    paragraph = document.add_heading(level=item["level"])
                else:
                    style = {
                        "bullet": "List Bullet",
                        "numbered": "List Number",
                    }.get(item["kind"])
                    paragraph = document.add_paragraph(style=style)
                run = paragraph.add_run(item["text"])
                run.bold = item["bold"]
                run.italic = item["italic"]
            document.save(temporary_path)
            return self._finalize(
                temporary_path,
                safe_name,
                "docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                {
                    "paragraphs": len(normalized),
                    "text_chars": total_chars,
                },
            )
        finally:
            temporary_path.unlink(missing_ok=True)

    def store_versioned_bytes(
        self,
        filename: str,
        data: bytes,
        artifact_format: str,
        mime_type: str,
        details: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Persist already-built artifact bytes as a NEW versioned file.

        Used by the editing subsystem: the original file is never overwritten,
        the change is always written next to it as ``name (n).ext``.
        """
        if not isinstance(data, (bytes, bytearray)) or len(data) == 0:
            raise ValueError("artifact bytes must be non-empty")
        safe_name = self._safe_filename(filename, artifact_format)
        temporary_path = self._temporary_path()
        temporary_path.write_bytes(bytes(data))
        try:
            return self._finalize(
                temporary_path,
                safe_name,
                artifact_format,
                mime_type,
                details or {},
            )
        finally:
            temporary_path.unlink(missing_ok=True)


__all__ = ["WorkspaceArtifactGenerator"]
