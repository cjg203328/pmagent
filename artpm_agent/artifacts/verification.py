"""Format-aware verification for generated workspace artifacts."""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any


class ArtifactVerificationError(RuntimeError):
    """Raised when a generated file cannot prove its delivery contract."""


def _require_file(path: Path) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise ArtifactVerificationError("generated artifact is empty or missing")


def _verify_xlsx(path: Path, expected: Mapping[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        expected_sheets = list(expected["sheets"])
        if workbook.sheetnames != [str(sheet["name"]) for sheet in expected_sheets]:
            raise ArtifactVerificationError("worksheet list does not match the plan")
        total_rows = 0
        max_columns = 0
        for sheet in expected_sheets:
            expected_values = [tuple(row) for row in sheet["values"]]
            expected_columns = len(expected_values[0]) if expected_values else 0
            values = list(
                workbook[str(sheet["name"])].iter_rows(
                    max_col=expected_columns,
                    values_only=True,
                )
            )
            if values != expected_values:
                raise ArtifactVerificationError(
                    "worksheet content does not match the plan"
                )
            total_rows += max(0, len(values) - 1)
            max_columns = max(max_columns, len(values[0]) if values else 0)
        return {
            "status": "passed",
            "reopened": True,
            "sheet_count": len(workbook.sheetnames),
            "row_count": total_rows,
            "column_count": max_columns,
            "content_match": True,
        }
    finally:
        workbook.close()


def _verify_docx(path: Path, expected: Mapping[str, Any]) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    actual = [paragraph.text for paragraph in document.paragraphs]
    expected_text = [str(value) for value in expected["paragraphs"]]
    if actual != expected_text:
        raise ArtifactVerificationError("document text does not match the plan")
    return {
        "status": "passed",
        "reopened": True,
        "paragraph_count": len(actual),
        "text_present": all(text in "\n".join(actual) for text in expected_text),
        "content_match": True,
    }


def _slide_text(slide: Any) -> str:
    return "\n".join(
        str(shape.text)
        for shape in slide.shapes
        if hasattr(shape, "text") and str(shape.text).strip()
    )


def _verify_pptx(path: Path, expected: Mapping[str, Any]) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as error:
        raise ArtifactVerificationError(
            "PPTX verification requires the documents profile"
        ) from error

    presentation = Presentation(path)
    expected_slides = list(expected["slides"])
    if len(presentation.slides) != len(expected_slides):
        raise ArtifactVerificationError(
            "presentation slide count does not match the plan"
        )
    actual_text = [_slide_text(slide) for slide in presentation.slides]
    for index, slide in enumerate(expected_slides):
        required = [str(slide["title"]), *map(str, slide.get("bullets", []))]
        if any(text not in actual_text[index] for text in required if text):
            raise ArtifactVerificationError(
                f"presentation slide {index + 1} is missing planned text"
            )
    return {
        "status": "passed",
        "reopened": True,
        "slide_count": len(actual_text),
        "text_present": True,
        "content_match": True,
    }


def _verify_pdf(path: Path, expected: Mapping[str, Any]) -> dict[str, Any]:
    try:
        import pymupdf
    except ImportError as error:
        raise ArtifactVerificationError(
            "PDF verification requires the documents profile"
        ) from error

    document = pymupdf.open(path)
    try:
        if document.page_count <= 0:
            raise ArtifactVerificationError("PDF contains no pages")
        extracted = "\n".join(page.get_text() for page in document)
        expected_text = [str(value) for value in expected["paragraphs"]]
        text_present = all(text in extracted for text in expected_text)
        if not text_present:
            raise ArtifactVerificationError("PDF text does not match the plan")
        return {
            "status": "passed",
            "reopened": True,
            "page_count": document.page_count,
            "text_present": True,
            "content_match": True,
        }
    finally:
        document.close()


def verify_artifact(
    path: str | Path,
    artifact_format: str,
    expected: Mapping[str, Any],
) -> dict[str, Any]:
    """Reopen a generated file and prove its planned content is present."""
    artifact_path = Path(path).resolve()
    _require_file(artifact_path)
    verifiers = {
        "xlsx": _verify_xlsx,
        "docx": _verify_docx,
        "pptx": _verify_pptx,
        "pdf": _verify_pdf,
    }
    normalized = str(artifact_format).lower().lstrip(".")
    verifier = verifiers.get(normalized)
    if verifier is None:
        raise ArtifactVerificationError(
            f"unsupported verification format: {normalized}"
        )
    result = verifier(artifact_path, expected)
    result["format"] = normalized
    result["size_bytes"] = artifact_path.stat().st_size
    return result


__all__ = ["ArtifactVerificationError", "verify_artifact"]
