"""Local spreadsheet template memory for generated artifacts.

Templates are saved only after an explicit user request. The store keeps
structure and provenance, not arbitrary conversation text, so the artifact
layer can reuse formats without polluting the broader knowledge base.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
import csv
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import tempfile
from typing import Any
import unicodedata
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook


MAX_TEMPLATE_NAME_CHARS = 80
MAX_TEMPLATE_SHEETS = 8
MAX_TEMPLATE_COLUMNS = 120
MAX_TEMPLATE_PARAGRAPHS = 80
MAX_HEADER_SCAN_ROWS = 20
MAX_SAMPLE_ROWS = 5
MAX_SEARCH_RESULTS = 20
STORE_VERSION = 1
SUPPORTED_TEMPLATE_EXTENSIONS = frozenset({"csv", "xlsx", "xlsm"})
SUPPORTED_DOCUMENT_TEMPLATE_EXTENSIONS = frozenset({"docx"})

_PATH_SEPARATORS = re.compile(r"[/\\\x00-\x1f]")
_TOKEN_RE = re.compile(r"[a-z0-9]+|[\u4e00-\u9fff]{1,6}", re.IGNORECASE)
_SPLIT_RE = re.compile(r"[\s,，、;；|/]+")


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def _clean_text(value: Any, *, max_chars: int | None = None) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"\s+", " ", text)
    if max_chars is not None:
        text = text[:max_chars].rstrip()
    return text


def _safe_template_name(value: Any, fallback: str) -> str:
    name = _clean_text(value, max_chars=MAX_TEMPLATE_NAME_CHARS)
    if not name:
        name = _clean_text(fallback, max_chars=MAX_TEMPLATE_NAME_CHARS)
    name = _PATH_SEPARATORS.sub("_", name).strip(" ._")
    return name or "表格模板"


def _dedupe(values: Iterable[str], *, max_items: int) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = _clean_text(value)
        key = text.casefold()
        if not text or key in seen:
            continue
        seen.add(key)
        result.append(text)
        if len(result) >= max_items:
            break
    return result


def _tokens(text: str) -> set[str]:
    normalized = _clean_text(text).casefold()
    tokens = {match.group(0) for match in _TOKEN_RE.finditer(normalized)}
    tokens.update(
        item
        for item in _SPLIT_RE.split(normalized)
        if item and len(item) <= MAX_TEMPLATE_NAME_CHARS
    )
    return tokens


def _jsonable_cell(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalize_columns(values: Sequence[Any]) -> list[str]:
    columns: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values[:MAX_TEMPLATE_COLUMNS], start=1):
        text = _clean_text(value, max_chars=120)
        if not text:
            text = f"列{index}"
        base = text
        key = base.casefold()
        if key in used:
            used[key] += 1
            text = f"{base}_{used[key]}"
        else:
            used[key] = 1
        columns.append(text)
    return columns


def _non_empty_count(values: Sequence[Any]) -> int:
    return sum(1 for value in values if _clean_text(value))


def _text_cell_count(values: Sequence[Any]) -> int:
    return sum(1 for value in values if isinstance(value, str) and value.strip())


def _row_to_mapping(columns: Sequence[str], values: Sequence[Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for index, column in enumerate(columns):
        value = values[index] if index < len(values) else None
        row[column] = _jsonable_cell(value)
    return row


def _sample_rows_from_rows(
    rows: Sequence[Sequence[Any]],
    *,
    header_index: int,
    columns: Sequence[str],
) -> list[dict[str, Any]]:
    samples: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        values = list(row[: len(columns)])
        if _non_empty_count(values) <= 0:
            continue
        samples.append(_row_to_mapping(columns, values))
        if len(samples) >= MAX_SAMPLE_ROWS:
            break
    return samples


def _detect_header(rows: Sequence[Sequence[Any]]) -> tuple[int, list[str]]:
    best_index = 0
    best_score = (-1, -1, 0)
    for index, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        trimmed = list(row[:MAX_TEMPLATE_COLUMNS])
        score = (_non_empty_count(trimmed), _text_cell_count(trimmed), -index)
        if score > best_score:
            best_score = score
            best_index = index
    raw_columns = list(rows[best_index][:MAX_TEMPLATE_COLUMNS]) if rows else []
    if _non_empty_count(raw_columns) <= 0:
        raw_columns = ["列1"]
    return best_index, _normalize_columns(raw_columns)


def _workbook_sheets(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets[:MAX_TEMPLATE_SHEETS]:
            rows = [
                tuple(row)
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=min(worksheet.max_row or 1, MAX_HEADER_SCAN_ROWS + MAX_SAMPLE_ROWS),
                    values_only=True,
                )
            ]
            header_index, columns = _detect_header(rows)
            samples = _sample_rows_from_rows(
                rows,
                header_index=header_index,
                columns=columns,
            )
            sheets.append(
                {
                    "name": _clean_text(worksheet.title, max_chars=31) or "Sheet1",
                    "header_row": header_index + 1,
                    "columns": columns,
                    "sample_rows": samples,
                    "row_count": int(worksheet.max_row or 0),
                    "column_count": int(worksheet.max_column or len(columns)),
                }
            )
        return sheets
    finally:
        workbook.close()


def _csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    return list(csv.reader(text.splitlines()))[: MAX_HEADER_SCAN_ROWS + MAX_SAMPLE_ROWS]


def _csv_sheets(path: Path) -> list[dict[str, Any]]:
    rows = _csv_rows(path)
    header_index, columns = _detect_header(rows)
    samples = _sample_rows_from_rows(
        rows,
        header_index=header_index,
        columns=columns,
    )
    return [
        {
            "name": _clean_text(path.stem, max_chars=31) or "Sheet1",
            "header_row": header_index + 1,
            "columns": columns,
            "sample_rows": samples,
            "row_count": len(rows),
            "column_count": len(columns),
        }
    ]


def extract_spreadsheet_template(
    file_path: str | Path,
    *,
    name: str | None = None,
    keywords: Iterable[str] | None = None,
    notes: str | None = None,
    source: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Extract bounded spreadsheet structure from one local file."""
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Template source does not exist: {path}")
    extension = path.suffix.lower().lstrip(".")
    if extension not in SUPPORTED_TEMPLATE_EXTENSIONS:
        raise ValueError(f"Unsupported spreadsheet template extension: .{extension}")

    if extension == "csv":
        sheets = _csv_sheets(path)
    else:
        sheets = _workbook_sheets(path)
    if not sheets:
        raise ValueError("Spreadsheet template has no readable worksheets")

    digest = sha256(path.read_bytes()).hexdigest()
    title = _safe_template_name(name, path.stem)
    keyword_values = _dedupe(
        [
            *(keywords or []),
            title,
            path.stem,
            *(sheet["name"] for sheet in sheets),
            *(column for sheet in sheets for column in sheet["columns"]),
        ],
        max_items=32,
    )
    now = _utc_now()
    return {
        "id": uuid4().hex,
        "kind": "spreadsheet",
        "name": title,
        "keywords": keyword_values,
        "notes": _clean_text(notes, max_chars=1000),
        "sheets": sheets,
        "source": dict(source or {}),
        "source_hash": digest,
        "source_name": path.name,
        "created_at": now,
        "updated_at": now,
        "last_used_at": None,
        "usage_count": 0,
    }


def _paragraph_kind(style_name: str) -> tuple[str, int]:
    normalized = style_name.casefold()
    if normalized.startswith("heading"):
        match = re.search(r"(\d+)", normalized)
        level = int(match.group(1)) if match else 1
        return "heading", max(1, min(level, 9))
    if "bullet" in normalized:
        return "bullet", 1
    if "number" in normalized or "list" in normalized:
        return "numbered", 1
    return "paragraph", 1


def _paragraph_flags(paragraph: Any) -> dict[str, bool]:
    runs = list(getattr(paragraph, "runs", []) or [])
    return {
        "bold": any(run.bold is True for run in runs),
        "italic": any(run.italic is True for run in runs),
    }


def _document_outline(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    outline: list[dict[str, Any]] = []
    for paragraph in document.paragraphs:
        text = _clean_text(paragraph.text, max_chars=600)
        if not text:
            continue
        kind, level = _paragraph_kind(str(paragraph.style.name or ""))
        flags = _paragraph_flags(paragraph)
        outline.append(
            {
                "kind": kind,
                "level": level,
                "text": text,
                "bold": flags["bold"],
                "italic": flags["italic"],
            }
        )
        if len(outline) >= MAX_TEMPLATE_PARAGRAPHS:
            break
    if not outline:
        outline.append(
            {
                "kind": "paragraph",
                "level": 1,
                "text": path.stem,
                "bold": False,
                "italic": False,
            }
        )
    return outline


def extract_document_template(
    file_path: str | Path,
    *,
    name: str | None = None,
    keywords: Iterable[str] | None = None,
    notes: str | None = None,
    source: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Extract bounded DOCX structure from one local file."""
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Template source does not exist: {path}")
    extension = path.suffix.lower().lstrip(".")
    if extension not in SUPPORTED_DOCUMENT_TEMPLATE_EXTENSIONS:
        raise ValueError(f"Unsupported document template extension: .{extension}")

    outline = _document_outline(path)
    digest = sha256(path.read_bytes()).hexdigest()
    title = _safe_template_name(name, path.stem)
    keyword_values = _dedupe(
        [
            *(keywords or []),
            title,
            path.stem,
            *(item["text"] for item in outline[:12]),
            *(item["kind"] for item in outline),
        ],
        max_items=40,
    )
    now = _utc_now()
    return {
        "id": uuid4().hex,
        "kind": "document",
        "name": title,
        "keywords": keyword_values,
        "notes": _clean_text(notes, max_chars=1000),
        "paragraphs": outline,
        "source": dict(source or {}),
        "source_hash": digest,
        "source_name": path.name,
        "created_at": now,
        "updated_at": now,
        "last_used_at": None,
        "usage_count": 0,
    }


class SpreadsheetTemplateStore:
    """Persist and rank reusable spreadsheet templates for one workspace."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path).expanduser().resolve()
        if self.path.suffix.lower() != ".json":
            self.path = self.path / ".memory" / "spreadsheet_templates.json"

    def _load(self) -> dict[str, Any]:
        if not self.path.exists():
            return {"version": STORE_VERSION, "templates": []}
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as error:
            raise ValueError("Spreadsheet template store is not valid JSON") from error
        if not isinstance(data, Mapping):
            raise ValueError("Spreadsheet template store root must be a mapping")
        templates = data.get("templates", [])
        if not isinstance(templates, list):
            raise ValueError("Spreadsheet template store templates must be a list")
        return {"version": int(data.get("version", STORE_VERSION)), "templates": templates}

    def _save(self, data: Mapping[str, Any]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        payload = json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8",
            dir=str(self.path.parent),
            delete=False,
            prefix=".spreadsheet-templates-",
            suffix=".tmp",
        ) as handle:
            handle.write(payload)
            temporary = Path(handle.name)
        try:
            temporary.replace(self.path)
        finally:
            temporary.unlink(missing_ok=True)

    def learn_from_file(
        self,
        file_path: str | Path,
        *,
        name: str | None = None,
        keywords: Iterable[str] | None = None,
        notes: str | None = None,
        source: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Save or update one template from an explicit user-provided file."""
        template = extract_spreadsheet_template(
            file_path,
            name=name,
            keywords=keywords,
            notes=notes,
            source=source,
        )
        data = self._load()
        templates = [dict(item) for item in data["templates"] if isinstance(item, Mapping)]

        name_key = template["name"].casefold()
        existing_index = next(
            (
                index
                for index, item in enumerate(templates)
                if str(item.get("name", "")).casefold() == name_key
                or item.get("source_hash") == template["source_hash"]
            ),
            None,
        )
        if existing_index is not None:
            existing = templates[existing_index]
            template["id"] = str(existing.get("id") or template["id"])
            template["created_at"] = str(existing.get("created_at") or template["created_at"])
            template["usage_count"] = int(existing.get("usage_count") or 0)
            template["last_used_at"] = existing.get("last_used_at")
            templates[existing_index] = template
        else:
            templates.append(template)

        data["templates"] = templates
        data["version"] = STORE_VERSION
        self._save(data)
        return dict(template)

    def list_templates(self, *, limit: int = 50) -> list[dict[str, Any]]:
        if isinstance(limit, bool) or not isinstance(limit, int) or limit <= 0:
            raise ValueError("limit must be a positive integer")
        templates = [dict(item) for item in self._load()["templates"] if isinstance(item, Mapping)]
        templates.sort(
            key=lambda item: (
                str(item.get("updated_at") or ""),
                int(item.get("usage_count") or 0),
                str(item.get("id") or ""),
            ),
            reverse=True,
        )
        return templates[:limit]

    def get(self, template_id_or_name: str) -> dict[str, Any] | None:
        key = _clean_text(template_id_or_name).casefold()
        if not key:
            return None
        for item in self.list_templates(limit=10_000):
            if str(item.get("id", "")).casefold() == key:
                return item
            if str(item.get("name", "")).casefold() == key:
                return item
        return None

    def search(self, query: str, *, limit: int = 5) -> list[dict[str, Any]]:
        if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= MAX_SEARCH_RESULTS:
            raise ValueError(f"limit must be between 1 and {MAX_SEARCH_RESULTS}")
        query_text = _clean_text(query).casefold()
        if not query_text:
            return []
        query_tokens = _tokens(query_text)
        scored: list[tuple[float, str, dict[str, Any]]] = []
        for template in self.list_templates(limit=10_000):
            haystack_parts = [
                str(template.get("name") or ""),
                str(template.get("notes") or ""),
                " ".join(str(item) for item in template.get("keywords") or []),
            ]
            for sheet in template.get("sheets") or []:
                if isinstance(sheet, Mapping):
                    haystack_parts.append(str(sheet.get("name") or ""))
                    haystack_parts.extend(str(item) for item in sheet.get("columns") or [])
            haystack = " ".join(haystack_parts).casefold()
            score = 0.0
            name = str(template.get("name") or "").casefold()
            if name and name in query_text:
                score += 60.0
            for keyword in template.get("keywords") or []:
                keyword_text = str(keyword).casefold().strip()
                if keyword_text and keyword_text in query_text:
                    score += 12.0
            haystack_tokens = _tokens(haystack)
            score += len(query_tokens & haystack_tokens) * 3.0
            if query_text and query_text in haystack:
                score += 8.0
            score += min(int(template.get("usage_count") or 0), 20) * 0.1
            if score <= 0:
                continue
            ranked = dict(template)
            ranked["score"] = score
            scored.append((score, str(template.get("updated_at") or ""), ranked))
        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
        return [item[2] for item in scored[:limit]]

    def record_use(self, template_id: str) -> dict[str, Any] | None:
        data = self._load()
        templates = [dict(item) for item in data["templates"] if isinstance(item, Mapping)]
        now = _utc_now()
        updated: dict[str, Any] | None = None
        for index, template in enumerate(templates):
            if str(template.get("id")) != template_id:
                continue
            template["usage_count"] = int(template.get("usage_count") or 0) + 1
            template["last_used_at"] = now
            template["updated_at"] = now
            templates[index] = template
            updated = dict(template)
            break
        if updated is not None:
            data["templates"] = templates
            data["version"] = STORE_VERSION
            self._save(data)
        return updated


class DocumentTemplateStore:
    """Persist and rank reusable DOCX templates for one workspace."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path).expanduser().resolve()
        if self.path.suffix.lower() != ".json":
            self.path = self.path / ".memory" / "document_templates.json"

    def _load(self) -> dict[str, Any]:
        if not self.path.exists():
            return {"version": STORE_VERSION, "templates": []}
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as error:
            raise ValueError("Document template store is not valid JSON") from error
        if not isinstance(data, Mapping):
            raise ValueError("Document template store root must be a mapping")
        templates = data.get("templates", [])
        if not isinstance(templates, list):
            raise ValueError("Document template store templates must be a list")
        return {"version": int(data.get("version", STORE_VERSION)), "templates": templates}

    def _save(self, data: Mapping[str, Any]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        payload = json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8",
            dir=str(self.path.parent),
            delete=False,
            prefix=".document-templates-",
            suffix=".tmp",
        ) as handle:
            handle.write(payload)
            temporary = Path(handle.name)
        try:
            temporary.replace(self.path)
        finally:
            temporary.unlink(missing_ok=True)

    def learn_from_file(
        self,
        file_path: str | Path,
        *,
        name: str | None = None,
        keywords: Iterable[str] | None = None,
        notes: str | None = None,
        source: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Save or update one document template from an explicit user file."""
        template = extract_document_template(
            file_path,
            name=name,
            keywords=keywords,
            notes=notes,
            source=source,
        )
        data = self._load()
        templates = [dict(item) for item in data["templates"] if isinstance(item, Mapping)]

        name_key = template["name"].casefold()
        existing_index = next(
            (
                index
                for index, item in enumerate(templates)
                if str(item.get("name", "")).casefold() == name_key
                or item.get("source_hash") == template["source_hash"]
            ),
            None,
        )
        if existing_index is not None:
            existing = templates[existing_index]
            template["id"] = str(existing.get("id") or template["id"])
            template["created_at"] = str(existing.get("created_at") or template["created_at"])
            template["usage_count"] = int(existing.get("usage_count") or 0)
            template["last_used_at"] = existing.get("last_used_at")
            templates[existing_index] = template
        else:
            templates.append(template)

        data["templates"] = templates
        data["version"] = STORE_VERSION
        self._save(data)
        return dict(template)

    def list_templates(self, *, limit: int = 50) -> list[dict[str, Any]]:
        if isinstance(limit, bool) or not isinstance(limit, int) or limit <= 0:
            raise ValueError("limit must be a positive integer")
        templates = [dict(item) for item in self._load()["templates"] if isinstance(item, Mapping)]
        templates.sort(
            key=lambda item: (
                str(item.get("updated_at") or ""),
                int(item.get("usage_count") or 0),
                str(item.get("id") or ""),
            ),
            reverse=True,
        )
        return templates[:limit]

    def get(self, template_id_or_name: str) -> dict[str, Any] | None:
        key = _clean_text(template_id_or_name).casefold()
        if not key:
            return None
        for item in self.list_templates(limit=10_000):
            if str(item.get("id", "")).casefold() == key:
                return item
            if str(item.get("name", "")).casefold() == key:
                return item
        return None

    def search(self, query: str, *, limit: int = 5) -> list[dict[str, Any]]:
        if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= MAX_SEARCH_RESULTS:
            raise ValueError(f"limit must be between 1 and {MAX_SEARCH_RESULTS}")
        query_text = _clean_text(query).casefold()
        if not query_text:
            return []
        query_tokens = _tokens(query_text)
        scored: list[tuple[float, str, dict[str, Any]]] = []
        for template in self.list_templates(limit=10_000):
            haystack_parts = [
                str(template.get("name") or ""),
                str(template.get("notes") or ""),
                " ".join(str(item) for item in template.get("keywords") or []),
            ]
            for paragraph in template.get("paragraphs") or []:
                if isinstance(paragraph, Mapping):
                    haystack_parts.append(str(paragraph.get("text") or ""))
                    haystack_parts.append(str(paragraph.get("kind") or ""))
            haystack = " ".join(haystack_parts).casefold()
            score = 0.0
            name = str(template.get("name") or "").casefold()
            if name and name in query_text:
                score += 60.0
            for keyword in template.get("keywords") or []:
                keyword_text = str(keyword).casefold().strip()
                if keyword_text and keyword_text in query_text:
                    score += 12.0
            haystack_tokens = _tokens(haystack)
            score += len(query_tokens & haystack_tokens) * 3.0
            if query_text and query_text in haystack:
                score += 8.0
            score += min(int(template.get("usage_count") or 0), 20) * 0.1
            if score <= 0:
                continue
            ranked = dict(template)
            ranked["score"] = score
            scored.append((score, str(template.get("updated_at") or ""), ranked))
        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
        return [item[2] for item in scored[:limit]]

    def record_use(self, template_id: str) -> dict[str, Any] | None:
        data = self._load()
        templates = [dict(item) for item in data["templates"] if isinstance(item, Mapping)]
        now = _utc_now()
        updated: dict[str, Any] | None = None
        for index, template in enumerate(templates):
            if str(template.get("id")) != template_id:
                continue
            template["usage_count"] = int(template.get("usage_count") or 0) + 1
            template["last_used_at"] = now
            template["updated_at"] = now
            templates[index] = template
            updated = dict(template)
            break
        if updated is not None:
            data["templates"] = templates
            data["version"] = STORE_VERSION
            self._save(data)
        return updated


def template_context(template: Mapping[str, Any], *, max_chars: int = 1800) -> str:
    """Return bounded text for a model planning prompt."""
    sheets = []
    for sheet in template.get("sheets") or []:
        if not isinstance(sheet, Mapping):
            continue
        sheets.append(
            {
                "name": sheet.get("name"),
                "columns": list(sheet.get("columns") or [])[:MAX_TEMPLATE_COLUMNS],
                "sample_rows": list(sheet.get("sample_rows") or [])[:2],
            }
        )
    payload = {
        "template_name": template.get("name"),
        "rules": "Keep sheet names and columns exactly. Fill only user-provided data.",
        "sheets": sheets,
    }
    return json.dumps(payload, ensure_ascii=False, default=str)[:max_chars]


def document_template_context(
    template: Mapping[str, Any],
    *,
    max_chars: int = 2200,
) -> str:
    """Return bounded document-template text for a model planning prompt."""
    paragraphs = []
    for paragraph in template.get("paragraphs") or []:
        if not isinstance(paragraph, Mapping):
            continue
        paragraphs.append(
            {
                "kind": paragraph.get("kind", "paragraph"),
                "level": paragraph.get("level", 1),
                "text_sample": paragraph.get("text", ""),
                "bold": bool(paragraph.get("bold", False)),
                "italic": bool(paragraph.get("italic", False)),
            }
        )
        if len(paragraphs) >= MAX_TEMPLATE_PARAGRAPHS:
            break
    payload = {
        "template_name": template.get("name"),
        "rules": (
            "Keep paragraph order and kind/level. Replace sample text with "
            "user-provided content. Do not invent paths or destructive actions."
        ),
        "paragraphs": paragraphs,
    }
    return json.dumps(payload, ensure_ascii=False, default=str)[:max_chars]


__all__ = [
    "DocumentTemplateStore",
    "SpreadsheetTemplateStore",
    "SUPPORTED_DOCUMENT_TEMPLATE_EXTENSIONS",
    "SUPPORTED_TEMPLATE_EXTENSIONS",
    "document_template_context",
    "extract_document_template",
    "extract_spreadsheet_template",
    "template_context",
]
