"""Preserve-format monthly filtering for uploaded Excel workbooks."""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class MonthlyXlsxError(ValueError):
    """Raised when a workbook cannot be filtered deterministically."""


@dataclass(frozen=True)
class MonthlyXlsxResult:
    target_month: str
    mode: str
    source_sha256: str
    source_name: str
    sheets: tuple[dict[str, Any], ...]
    matched_rows: int
    uncovered_rows: int
    month_column: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "target_month": self.target_month,
            "mode": self.mode,
            "source_sha256": self.source_sha256,
            "source_name": self.source_name,
            "sheets": [dict(sheet) for sheet in self.sheets],
            "matched_rows": self.matched_rows,
            "uncovered_rows": self.uncovered_rows,
            "month_column": self.month_column,
            "note": (
                f"口径：{self.target_month} / "
                f"{'截至目标月累计' if self.mode == 'cumulative' else '目标月新增'}；"
                f"保留 {self.matched_rows} 行，未覆盖 {self.uncovered_rows} 行。"
            ),
        }


_MONTH_HEADER = re.compile(r"月份|月度|month|日期|时间|开始|完成|交付", re.I)
_MONTH_VALUE = re.compile(
    r"(?P<year>20\d{2})\s*(?:年|[-/.])\s*(?P<month>1[0-2]|0?[1-9])\s*(?:月)?"
)
_MONTH_ONLY = re.compile(r"(?<!\d)(?P<month>1[0-2]|0?[1-9])\s*月(?!\d)")
_MODES = {"new_items", "cumulative"}


def normalize_target_month(value: str | date | datetime) -> str:
    """Normalize a target month to ``YYYY-MM``."""
    if isinstance(value, datetime | date):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value or "").strip()
    match = _MONTH_VALUE.search(text)
    if not match:
        raise MonthlyXlsxError("target_month must look like YYYY-MM or YYYY年M月")
    return f"{int(match.group('year')):04d}-{int(match.group('month')):02d}"


def normalize_cell_month(value: Any, *, default_year: int | None = None) -> str | None:
    """Normalize common Excel date/month cell values."""
    if isinstance(value, datetime | date):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, (int, float)):
        return None
    text = str(value or "").strip()
    if not text:
        return None
    match = _MONTH_VALUE.search(text)
    if match:
        return f"{int(match.group('year')):04d}-{int(match.group('month')):02d}"
    match = _MONTH_ONLY.search(text)
    if match and default_year is not None:
        return f"{default_year:04d}-{int(match.group('month')):02d}"
    return None


_MONTH_ACTION = re.compile(
    r"(?:只保留|仅保留|只留|只取|留下|保留|仅|汇总|裁剪|筛选|导出|另存|拆出|累计|结转)"
)
_CUMULATIVE_HINT = re.compile(
    r"(?:累计|结转|含之前|包括之前|包含之前|含上月|之前所有|截至|截止到)"
)


def is_monthly_request(prompt: str) -> bool:
    """True when the prompt asks to keep one month's rows out of a workbook."""
    text = str(prompt or "")
    has_month = bool(_MONTH_VALUE.search(text) or _MONTH_ONLY.search(text))
    return has_month and bool(_MONTH_ACTION.search(text))


def mode_from_prompt(prompt: str) -> str:
    """Cumulative keeps earlier months too; the default keeps only the target."""
    return "cumulative" if _CUMULATIVE_HINT.search(str(prompt or "")) else "new_items"


def _explicit_years(source_path: Path, month_number: int) -> tuple[int, ...]:
    """Collect the years the workbook itself uses for that month.

    Monthly workbooks record the month either in the sheet title (``24年11月``) or
    in a per-row month column, so both are scanned.
    """
    years: set[int] = set()
    workbook = load_workbook(source_path, data_only=True)
    try:
        for sheet in workbook.worksheets:
            titled = sheet_target_month(sheet.title)
            if titled is not None:
                if int(titled[5:]) == month_number:
                    years.add(int(titled[:4]))
                continue
            detection = _header_row_and_month_column(sheet)
            if detection is None:
                continue
            header_row, month_column, _header = detection
            for row_index in range(header_row + 1, sheet.max_row + 1):
                normalized = normalize_cell_month(
                    sheet.cell(row_index, month_column).value
                )
                if normalized and int(normalized[5:]) == month_number:
                    years.add(int(normalized[:4]))
    finally:
        workbook.close()
    return tuple(sorted(years))


def resolve_target_month(prompt: str, source_path: str | Path | None = None) -> str:
    """Resolve a prompt month to ``YYYY-MM`` without guessing a year.

    A bare month such as "11月" is only safe when the source workbook uses it in
    exactly one year. Otherwise the caller must ask: silently choosing the current
    year would filter to an empty sheet and still report success.
    """
    text = str(prompt or "")
    match = _MONTH_VALUE.search(text)
    if match:
        return f"{int(match.group('year')):04d}-{int(match.group('month')):02d}"
    month_only = _MONTH_ONLY.search(text)
    if month_only is None:
        raise MonthlyXlsxError("未识别到月份")
    month_number = int(month_only.group("month"))
    if source_path is None:
        raise MonthlyXlsxError(f"需要确认「{month_number}月」属于哪一年")
    years = _explicit_years(Path(source_path), month_number)
    if len(years) == 1:
        return f"{years[0]:04d}-{month_number:02d}"
    if not years:
        raise MonthlyXlsxError(f"表格中没有找到 {month_number}月 的数据")
    joined = "、".join(str(year) for year in years)
    raise MonthlyXlsxError(f"表格中 {month_number}月 同时出现在 {joined} 年，请指定年份")


_SHEET_SHORT_MONTH = re.compile(
    r"(?P<year>\d{2})\s*(?:年|[-/.])\s*(?P<month>1[0-2]|0?[1-9])\s*月?"
)


def sheet_target_month(title: Any) -> str | None:
    """Return ``YYYY-MM`` when a sheet title is itself a month label.

    Monthly capacity workbooks carry one sheet per month (``24年11月``), so the
    month lives in the title rather than in a column. Two-digit years are read as
    ``20YY``; anything that is not exactly a month label returns ``None``.
    """
    text = str(title or "").strip()
    if not text:
        return None
    full = _MONTH_VALUE.fullmatch(text)
    if full:
        return f"{int(full.group('year')):04d}-{int(full.group('month')):02d}"
    short = _SHEET_SHORT_MONTH.fullmatch(text)
    if short:
        return f"20{int(short.group('year')):02d}-{int(short.group('month')):02d}"
    return None


def _text_header_row(sheet: Any) -> int:
    """Best-effort header row for a wide sheet: first text-heavy row on top."""
    for row_index in range(1, min(sheet.max_row, 5) + 1):
        texts = [
            value
            for value in (
                sheet.cell(row_index, column).value
                for column in range(1, sheet.max_column + 1)
            )
            if isinstance(value, str) and value.strip()
        ]
        if len(texts) >= 3:
            return row_index
    return 1


def _select_sheets_by_month(
    workbook: Any,
    normalized_target: str,
    mode: str,
    titled: dict[str, str | None],
) -> tuple[list[dict[str, Any]], int, int, str]:
    """Wide-format workbooks: one sheet per month, selected by sheet title.

    Sheets whose title is not a month are kept. They are usually reference or
    summary tabs, and deleting data the tool cannot classify is the worse error.
    """
    for sheet in list(workbook.worksheets):
        month = titled.get(sheet.title)
        if month is None:
            continue
        include = month <= normalized_target if mode == "cumulative" else (
            month == normalized_target
        )
        if not include:
            del workbook[sheet.title]
    summaries: list[dict[str, Any]] = []
    total_matched = 0
    for sheet in workbook.worksheets:
        if sheet.title == "处理说明":
            continue
        header_row = _text_header_row(sheet)
        data_rows = max(0, sheet.max_row - header_row)
        total_matched += data_rows
        summaries.append(
            {
                "name": sheet.title,
                "header_row": header_row,
                "month_column": "工作表名称",
                "matched_rows": data_rows,
                "uncovered_rows": 0,
            }
        )
    if not any(titled.get(item["name"]) for item in summaries):
        raise MonthlyXlsxError(f"工作簿中没有 {normalized_target} 对应的工作表")
    return summaries, total_matched, 0, "工作表名称"


def _filter_rows_by_month(
    workbook: Any,
    normalized_target: str,
    mode: str,
    target_year: int,
) -> tuple[list[dict[str, Any]], int, int, str]:
    """Long-format tables: one row per record, filtered on a month column."""
    summaries: list[dict[str, Any]] = []
    total_matched = 0
    total_uncovered = 0
    detected_column = ""
    for sheet in workbook.worksheets:
        detection = _header_row_and_month_column(sheet)
        if detection is None:
            continue
        header_row, month_column, month_header = detection
        detected_column = month_header
        matched_rows: list[int] = []
        candidate_rows = 0
        for row_index in range(header_row + 1, sheet.max_row + 1):
            row_values = [
                sheet.cell(row_index, col).value
                for col in range(1, sheet.max_column + 1)
            ]
            if not any(value not in (None, "") for value in row_values):
                continue
            month = normalize_cell_month(
                sheet.cell(row_index, month_column).value, default_year=target_year
            )
            if month is None:
                continue
            candidate_rows += 1
            keep = month == normalized_target
            if mode == "cumulative":
                keep = month <= normalized_target
            if keep:
                matched_rows.append(row_index)
        candidate_row_set = {
            row_index
            for row_index in range(header_row + 1, sheet.max_row + 1)
            if normalize_cell_month(
                sheet.cell(row_index, month_column).value,
                default_year=target_year,
            )
            is not None
        }
        rows_to_delete = [
            row for row in candidate_row_set if row not in matched_rows
        ]
        for row_index in reversed(rows_to_delete):
            sheet.delete_rows(row_index, 1)
        total_matched += len(matched_rows)
        total_uncovered += max(0, candidate_rows - len(matched_rows))
        summaries.append(
            {
                "name": sheet.title,
                "header_row": header_row,
                "month_column": month_header,
                "matched_rows": len(matched_rows),
                "uncovered_rows": max(0, candidate_rows - len(matched_rows)),
            }
        )
    return summaries, total_matched, total_uncovered, detected_column


def _write_processing_note(
    workbook: Any,
    *,
    target: str,
    mode: str,
    column: str,
    matched: int,
    uncovered: int,
    digest: str,
) -> None:
    """Record the processing basis so the verifier can re-check the artifact."""
    note_sheet = workbook.create_sheet("处理说明")
    note_sheet.append(["字段", "值"])
    note_sheet.append(["目标月份", target])
    note_sheet.append(
        ["处理口径", "截至目标月累计" if mode == "cumulative" else "目标月新增"]
    )
    note_sheet.append(["月份列", column])
    note_sheet.append(["保留行数", matched])
    note_sheet.append(["未覆盖行数", uncovered])
    note_sheet.append(["源文件 SHA-256", digest])
    note_sheet.append(["说明", f"保留 {matched} 行，未覆盖 {uncovered} 行。"])
    note_sheet.freeze_panes = "A2"
    note_sheet.column_dimensions["A"].width = 18
    note_sheet.column_dimensions["B"].width = 80


def _header_row_and_month_column(sheet: Any) -> tuple[int, int, str] | None:
    candidates: list[tuple[int, int, str, int]] = []
    for row_index in range(1, min(sheet.max_row, 30) + 1):
        for column_index in range(1, sheet.max_column + 1):
            value = str(sheet.cell(row_index, column_index).value or "").strip()
            if _MONTH_HEADER.search(value):
                candidates.append((row_index, column_index, value, 3))
    for row_index, column_index, value, _score in candidates:
        for data_row in range(row_index + 1, min(sheet.max_row, row_index + 10) + 1):
            if normalize_cell_month(sheet.cell(data_row, column_index).value):
                return row_index, column_index, value
    return None


def _copy_row_style(sheet: Any, source_row: int, target_row: int) -> None:
    for source_cell, target_cell in zip(
        sheet[source_row],
        sheet[target_row],
    ):
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.hyperlink:
            target_cell._hyperlink = copy(source_cell.hyperlink)
        if source_cell.comment:
            target_cell.comment = copy(source_cell.comment)


def transform_monthly_workbook(
    source_path: str | Path,
    output_path: str | Path,
    target_month: str | date | datetime,
    *,
    mode: str = "new_items",
) -> MonthlyXlsxResult:
    """Copy a workbook, retain its formatting, and filter rows by month.

    Only ``.xlsx``/``.xlsm`` are supported because legacy ``.xls`` cannot be
    copied with openpyxl without losing workbook-level formatting.
    """
    source = Path(source_path).expanduser().resolve()
    destination = Path(output_path).expanduser().resolve()
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise MonthlyXlsxError(
            "preserve-format monthly filtering requires .xlsx or .xlsm"
        )
    if not source.is_file():
        raise MonthlyXlsxError("source workbook does not exist")
    if mode not in _MODES:
        raise MonthlyXlsxError("mode must be new_items or cumulative")
    normalized_target = normalize_target_month(target_month)
    target_year = int(normalized_target[:4])
    destination.parent.mkdir(parents=True, exist_ok=True)
    source_digest = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        titled = {
            sheet.title: sheet_target_month(sheet.title)
            for sheet in workbook.worksheets
        }
        if any(titled.values()):
            summaries, total_matched, total_uncovered, detected_column = (
                _select_sheets_by_month(workbook, normalized_target, mode, titled)
            )
        else:
            summaries, total_matched, total_uncovered, detected_column = (
                _filter_rows_by_month(
                    workbook, normalized_target, mode, target_year
                )
            )
        if not summaries:
            raise MonthlyXlsxError(
                "could not identify month-named sheets or a month column with "
                "date-like values"
            )
        _write_processing_note(
            workbook,
            target=normalized_target,
            mode=mode,
            column=detected_column,
            matched=total_matched,
            uncovered=total_uncovered,
            digest=source_digest,
        )
        workbook.save(destination)
    finally:
        workbook.close()
    return MonthlyXlsxResult(
        target_month=normalized_target,
        mode=mode,
        source_sha256=source_digest,
        source_name=source.name,
        sheets=tuple(summaries),
        matched_rows=total_matched,
        uncovered_rows=total_uncovered,
        month_column=detected_column,
    )


__all__ = [
    "MonthlyXlsxError",
    "MonthlyXlsxResult",
    "is_monthly_request",
    "mode_from_prompt",
    "normalize_cell_month",
    "normalize_target_month",
    "resolve_target_month",
    "sheet_target_month",
    "transform_monthly_workbook",
]
