"""Tolerant workbook loading for spreadsheets written by non-strict tools.

Some exporters emit OOXML that Excel and WPS accept but openpyxl rejects, for
example an empty ``<fill/>`` element inside ``xl/styles.xml``. openpyxl raises
``TypeError: expected <class 'openpyxl.styles.fills.Fill'>`` while reading the
style table, which fails the whole workbook even though the cell data is fine.

The repair below rewrites only those malformed style nodes into the minimal
valid form and retries once, in memory. The caller's file is never modified and
no temporary file is created, because a read-only workbook keeps its handle
open and Windows refuses to unlink an open file.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, BinaryIO, Union

_EMPTY_FILL = re.compile(rb"<fill(\s[^>]*)?/>|<fill(\s[^>]*)?>\s*</fill>")
_REPLACEMENT = b'<fill><patternFill patternType="none"/></fill>'
_ZIP_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
_STYLES_MEMBER = "xl/styles.xml"


def repair_workbook_styles(raw: bytes) -> tuple[bytes, bool]:
    """Return archive bytes with malformed empty fill nodes made valid."""

    import zipfile

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if _STYLES_MEMBER not in archive.namelist():
            return raw, False
        styles = archive.read(_STYLES_MEMBER)
        repaired, count = _EMPTY_FILL.subn(_REPLACEMENT, styles)
        if not count:
            return raw, False
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for item in archive.infolist():
                payload = (
                    repaired if item.filename == _STYLES_MEMBER else archive.read(item)
                )
                target.writestr(item, payload)
    return buffer.getvalue(), True


def _suffix_of(source: Union[str, Path, BinaryIO], source_name: str) -> str:
    if isinstance(source, (str, Path)):
        return Path(source).suffix.lower()
    return Path(source_name).suffix.lower()


def _raw_bytes(source: Union[str, Path, BinaryIO], source_name: str) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if hasattr(source, "seek"):
        source.seek(0)
    data = source.read()
    if isinstance(data, str):
        raise TypeError("workbook stream must be opened in binary mode")
    return bytes(data)


def safe_load_workbook(
    source: Union[str, Path, BinaryIO],
    *,
    source_name: str = "",
    **kwargs: Any,
) -> Any:
    """Load a workbook, repairing tolerated-but-malformed style tables once.

    The original failure is re-raised when the repair does not apply or does not
    help, so callers never lose the real error cause.
    """

    import openpyxl

    if isinstance(source, (str, Path)):
        path = Path(source)
        target: Any = path
    else:
        if hasattr(source, "seek"):
            source.seek(0)
        target = source

    try:
        return openpyxl.load_workbook(target, **kwargs)
    except Exception as original_error:
        suffix = _suffix_of(source, source_name)
        if suffix not in _ZIP_SUFFIXES:
            raise
        try:
            repaired, changed = repair_workbook_styles(_raw_bytes(source, source_name))
        except Exception:
            raise original_error from None
        if not changed:
            raise original_error from None
        try:
            return openpyxl.load_workbook(io.BytesIO(repaired), **kwargs)
        except Exception:
            raise original_error from None


__all__ = ["safe_load_workbook", "repair_workbook_styles"]
