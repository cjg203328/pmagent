"""Deterministic local conversion of multimodal attachments to Markdown.

This module follows the same product idea as Microsoft MarkItDown: turn common
business files into compact Markdown that is easier for LLMs and retrieval
systems to consume. It intentionally uses narrow, local readers and never
executes embedded file content.
"""

from __future__ import annotations

import csv
import json
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED_MARKDOWN_SUFFIXES = frozenset(
    {
        ".csv",
        ".docx",
        ".json",
        ".md",
        ".pdf",
        ".pptx",
        ".txt",
        ".xls",
        ".xlsm",
        ".xlsx",
        ".jpeg",
        ".jpg",
        ".png",
        ".webp",
    }
)


@dataclass(frozen=True)
class MarkdownConversionResult:
    """A bounded Markdown representation of one attachment."""

    success: bool
    markdown: str = ""
    source_format: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)
    requires_vision: bool = False
    ocr_available: bool = False
    ocr_status: str = "not_applicable"
    truncated: bool = False
    error: str | None = None


class LocalMarkdownConverter:
    """Convert local attachments into compact Markdown for model context."""

    def __init__(
        self,
        *,
        max_chars: int = 32768,
        max_table_rows: int = 200,
        max_table_columns: int = 80,
        max_sheets: int = 8,
    ) -> None:
        self.max_chars = self._bounded_int(max_chars, "max_chars", 1024, 1_000_000)
        self.max_table_rows = self._bounded_int(
            max_table_rows,
            "max_table_rows",
            1,
            10_000,
        )
        self.max_table_columns = self._bounded_int(
            max_table_columns,
            "max_table_columns",
            1,
            500,
        )
        self.max_sheets = self._bounded_int(max_sheets, "max_sheets", 1, 100)

    @staticmethod
    def _bounded_int(value: int, field: str, minimum: int, maximum: int) -> int:
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError(f"{field} must be an integer")
        if not minimum <= value <= maximum:
            raise ValueError(f"{field} must be between {minimum} and {maximum}")
        return value

    def convert(
        self,
        file_path: str | Path,
        *,
        parsed_result: Mapping[str, Any] | None = None,
    ) -> MarkdownConversionResult:
        """Convert a supported local file to Markdown without side effects."""
        path = Path(file_path).expanduser().resolve()
        if not path.is_file():
            return MarkdownConversionResult(
                success=False,
                source_format=path.suffix.lower(),
                error=f"file does not exist: {path}",
            )

        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_MARKDOWN_SUFFIXES:
            return MarkdownConversionResult(
                success=False,
                source_format=suffix,
                error=f"unsupported markdown conversion format: {suffix or 'none'}",
            )

        try:
            if suffix in {".xlsx", ".xlsm"}:
                return self._convert_xlsx(path)
            if suffix == ".xls":
                return self._convert_xls(path)
            if suffix == ".csv":
                return self._convert_csv(path)
            if suffix == ".docx":
                return self._convert_docx(path)
            if suffix == ".pdf":
                return self._convert_pdf(path, parsed_result=parsed_result)
            if suffix == ".pptx":
                return self._convert_pptx(path)
            if suffix in {".txt", ".md"}:
                return self._convert_text(path, markdown=suffix == ".md")
            if suffix == ".json":
                return self._convert_json(path)
            if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
                return self._convert_image(path, parsed_result=parsed_result)
        except Exception as error:
            return MarkdownConversionResult(
                success=False,
                source_format=suffix,
                error=str(error),
            )

        return MarkdownConversionResult(
            success=False,
            source_format=suffix,
            error=f"unhandled markdown conversion format: {suffix}",
        )

    def _finish(
        self,
        markdown: str,
        *,
        source_format: str,
        metadata: Mapping[str, Any] | None = None,
        requires_vision: bool = False,
        ocr_available: bool = False,
        ocr_status: str = "not_applicable",
    ) -> MarkdownConversionResult:
        text = str(markdown or "").strip()
        truncated = False
        if len(text) > self.max_chars:
            text = text[: self.max_chars].rstrip() + "\n\n<!-- truncated -->"
            truncated = True
        return MarkdownConversionResult(
            success=bool(text),
            markdown=text,
            source_format=source_format,
            metadata=dict(metadata or {}),
            requires_vision=requires_vision,
            ocr_available=ocr_available,
            ocr_status=ocr_status,
            truncated=truncated,
            error=None if text else "no markdown content extracted",
        )

    @staticmethod
    def _title(path: Path) -> str:
        return f"# {path.name}"

    @staticmethod
    def _cell(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).replace("\r", "\n").replace("\n", "<br>")
        return text.replace("|", "\\|").strip()

    @staticmethod
    def _non_empty(values: Sequence[Any]) -> int:
        return sum(1 for value in values if str(value or "").strip())

    def _table_markdown(self, rows: Sequence[Sequence[Any]]) -> str:
        bounded_rows = [list(row[: self.max_table_columns]) for row in rows]
        bounded_rows = [row for row in bounded_rows if self._non_empty(row)]
        if not bounded_rows:
            return ""

        header_index = 0
        for index, row in enumerate(bounded_rows[:20]):
            if self._non_empty(row) > self._non_empty(bounded_rows[header_index]):
                header_index = index

        headers = [
            self._cell(value) or f"Column {idx + 1}"
            for idx, value in enumerate(bounded_rows[header_index])
        ]
        body_rows = bounded_rows[
            header_index + 1 : header_index + 1 + self.max_table_rows
        ]
        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
        ]
        for row in body_rows:
            padded = row + [""] * max(0, len(headers) - len(row))
            lines.append(
                "| "
                + " | ".join(self._cell(value) for value in padded[: len(headers)])
                + " |"
            )
        return "\n".join(lines)

    def _convert_xlsx(self, path: Path) -> MarkdownConversionResult:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sections = [self._title(path)]
            sheet_count = 0
            for worksheet in workbook.worksheets[: self.max_sheets]:
                rows = [
                    tuple(row)
                    for row in worksheet.iter_rows(
                        min_row=1,
                        max_row=self.max_table_rows + 20,
                        max_col=self.max_table_columns,
                        values_only=True,
                    )
                ]
                table = self._table_markdown(rows)
                if not table:
                    continue
                sheet_count += 1
                sections.extend([f"## Sheet: {worksheet.title}", table])
            return self._finish(
                "\n\n".join(sections),
                source_format=path.suffix.lower(),
                metadata={"sheet_count": sheet_count},
            )
        finally:
            workbook.close()

    def _convert_xls(self, path: Path) -> MarkdownConversionResult:
        import pandas as pd

        sheets = pd.read_excel(path, sheet_name=None, header=None)
        sections = [self._title(path)]
        for sheet_index, (sheet_name, frame) in enumerate(sheets.items()):
            if sheet_index >= self.max_sheets:
                break
            rows = frame.fillna("").values.tolist()[: self.max_table_rows + 20]
            table = self._table_markdown(rows)
            if table:
                sections.extend([f"## Sheet: {sheet_name}", table])
        return self._finish(
            "\n\n".join(sections),
            source_format=path.suffix.lower(),
            metadata={"sheet_count": min(len(sheets), self.max_sheets)},
        )

    def _convert_csv(self, path: Path) -> MarkdownConversionResult:
        rows = self._read_csv_rows(path)
        return self._finish(
            "\n\n".join([self._title(path), self._table_markdown(rows)]),
            source_format=".csv",
            metadata={"row_count": max(0, len(rows) - 1)},
        )

    def _read_csv_rows(self, path: Path) -> list[list[str]]:
        raw = path.read_bytes()
        text = ""
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = raw.decode("utf-8", errors="replace")
        return list(csv.reader(text.splitlines()))[: self.max_table_rows + 20]

    def _convert_docx(self, path: Path) -> MarkdownConversionResult:
        from docx import Document

        document = Document(path)
        sections = [self._title(path)]
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = str(paragraph.style.name or "").casefold()
            if style_name.startswith("heading"):
                level = 2
                for token in style_name.split():
                    if token.isdigit():
                        level = max(2, min(int(token) + 1, 6))
                        break
                sections.append(f"{'#' * level} {text}")
            elif "bullet" in style_name:
                sections.append(f"- {text}")
            elif "number" in style_name:
                sections.append(f"1. {text}")
            else:
                sections.append(text)
        for table_index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            markdown_table = self._table_markdown(rows)
            if markdown_table:
                sections.extend([f"## Table {table_index}", markdown_table])
        return self._finish(
            "\n\n".join(sections),
            source_format=".docx",
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
        )

    def _convert_pptx(self, path: Path) -> MarkdownConversionResult:
        from pptx import Presentation

        presentation = Presentation(path)
        sections = [self._title(path)]
        for index, slide in enumerate(presentation.slides, start=1):
            lines = [
                str(shape.text).strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and str(shape.text).strip()
            ]
            if lines:
                sections.extend([f"## Slide {index}", "\n\n".join(lines)])
        return self._finish(
            "\n\n".join(sections),
            source_format=".pptx",
            metadata={"slide_count": len(presentation.slides)},
        )

    def _convert_pdf(
        self,
        path: Path,
        *,
        parsed_result: Mapping[str, Any] | None,
    ) -> MarkdownConversionResult:
        import pdfplumber

        page_texts: list[str] = []
        page_count = 0
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for index, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    page_texts.append(f"## Page {index}\n\n{text}")

        raw_text = str((parsed_result or {}).get("raw_text", "") or "").strip()
        if raw_text and raw_text not in "\n\n".join(page_texts):
            page_texts.append(f"## OCR Text\n\n{raw_text}")
        extracted = (parsed_result or {}).get("extracted_data")
        pages_requiring_ocr = []
        if isinstance(extracted, Mapping):
            pages_requiring_ocr = list(extracted.get("pages_requiring_ocr") or [])
        return self._finish(
            "\n\n".join([self._title(path), *page_texts]),
            source_format=".pdf",
            metadata={
                "page_count": page_count,
                "pages_requiring_ocr": pages_requiring_ocr,
            },
            requires_vision=bool(pages_requiring_ocr)
            and (parsed_result or {}).get("ocr_status") != "completed",
            ocr_available=(parsed_result or {}).get("ocr_status") == "completed",
            ocr_status=str((parsed_result or {}).get("ocr_status") or "not_needed"),
        )

    def _convert_text(self, path: Path, *, markdown: bool) -> MarkdownConversionResult:
        text = path.read_text(encoding="utf-8", errors="replace")
        body = text if markdown else f"{self._title(path)}\n\n{text}"
        return self._finish(
            body,
            source_format=".md" if markdown else ".txt",
            metadata={"chars": len(text)},
        )

    def _convert_json(self, path: Path) -> MarkdownConversionResult:
        raw = path.read_text(encoding="utf-8", errors="replace")
        try:
            payload = json.loads(raw)
            formatted = json.dumps(payload, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            formatted = raw
        return self._finish(
            f"{self._title(path)}\n\n```json\n{formatted}\n```",
            source_format=".json",
            metadata={"chars": len(raw)},
        )

    def _convert_image(
        self,
        path: Path,
        *,
        parsed_result: Mapping[str, Any] | None,
    ) -> MarkdownConversionResult:
        parsed = parsed_result or {}
        extracted = parsed.get("extracted_data")
        metadata = dict(extracted) if isinstance(extracted, Mapping) else {}
        raw_text = str(parsed.get("raw_text", "") or "").strip()
        sections = [
            self._title(path),
            "## Image Metadata",
            json.dumps(metadata, ensure_ascii=False, default=str),
        ]
        if raw_text:
            sections.extend(["## OCR Text", raw_text])
        else:
            sections.append("## OCR Text\n\nNo OCR text extracted.")
        return self._finish(
            "\n\n".join(sections),
            source_format=path.suffix.lower(),
            metadata=metadata,
            requires_vision=bool(parsed.get("requires_vision", True)),
            ocr_available=bool(parsed.get("ocr_available", False)),
            ocr_status=str(parsed.get("ocr_status") or "unavailable"),
        )


__all__ = [
    "SUPPORTED_MARKDOWN_SUFFIXES",
    "LocalMarkdownConverter",
    "MarkdownConversionResult",
]
