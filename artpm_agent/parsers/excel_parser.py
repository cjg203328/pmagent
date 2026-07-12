"""
真实的Excel解析器 - 智能识别报价单格式
"""
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet
from typing import BinaryIO, Dict, List, Optional, Any, Union
from pathlib import Path
import re
from datetime import datetime
import sys

# 添加父目录到路径
sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    from utils.logger import get_logger
    logger = get_logger(__name__)
except ImportError:
    import logging
    logger = logging.getLogger(__name__)


class ExcelQuoteParser:
    """Excel报价单解析器"""

    def __init__(self):
        self.client_templates = self._load_templates()

    def _load_templates(self) -> Dict:
        """加载客户模板规则"""
        return {
            "腾讯": {
                "keywords": ["腾讯", "tencent"],
                "header_row": 1,
                "columns": {
                    "asset_name": ["资产名称", "项目名称", "内容"],
                    "quantity": ["数量", "个数"],
                    "unit_price": ["单价", "价格"],
                    "total": ["总价", "小计", "合计"]
                }
            },
            "网易": {
                "keywords": ["网易", "netease"],
                "header_row": 2,
                "columns": {
                    "asset_name": ["制作内容", "资产"],
                    "quantity": ["数量"],
                    "unit_price": ["单价"],
                    "total": ["金额"]
                }
            },
            "米哈游": {
                "keywords": ["米哈游", "mihoyo"],
                "header_row": 1,
                "columns": {
                    "asset_name": ["资产名称"],
                    "quantity": ["数量"],
                    "unit_price": ["报价"],
                    "total": ["总计"]
                }
            }
        }

    def parse(
        self,
        file_path: Union[str, Path, BinaryIO],
        client_hint: str = None,
    ) -> Dict:
        """解析Excel报价单"""
        source_name = self._source_name(file_path)
        logger.info(f"开始解析Excel文件: {source_name}")
        workbook = None

        try:
            # 检查文件是否存在
            if isinstance(file_path, (str, Path)) and not Path(file_path).exists():
                logger.error(f"文件不存在: {file_path}")
                return {
                    "success": False,
                    "error": f"文件不存在: {file_path}",
                    "error_type": "FileNotFoundError"
                }

            # 加载工作簿
            logger.debug("加载Excel工作簿...")
            workbook = self._load_workbook(file_path, source_name)
            sheet = workbook.active
            logger.debug(f"工作表: {sheet.title}, 行数: {sheet.max_row}, 列数: {sheet.max_column}")

            # 识别客户和模板
            client = self._detect_client(sheet, client_hint)
            logger.info(f"识别客户: {client}")
            template = self.client_templates.get(client, self.client_templates["腾讯"])

            # 提取基本信息
            logger.debug("提取元数据...")
            metadata = self._extract_metadata(sheet, client)
            logger.debug(f"项目名称: {metadata.get('project_name')}")

            # 找到表头行
            header_row = self._find_header_row(sheet, template)
            logger.debug(f"表头行: {header_row}")

            # 解析列映射
            column_map = self._map_columns(sheet, header_row, template)
            logger.debug(f"列映射: {column_map}")

            # 提取资产数据
            logger.debug("提取资产数据...")
            assets = self._extract_assets(sheet, header_row + 1, column_map)
            logger.info(f"提取到 {len(assets)} 个资产")

            # 计算总金额
            total_amount = self._calculate_total(assets, sheet)
            logger.info(f"总金额: ¥{total_amount:,.2f}")

            result = {
                "success": True,
                "document_type": "报价单",
                "client": client,
                "project_name": metadata.get("project_name"),
                "date": metadata.get("date"),
                "contact": metadata.get("contact"),
                "assets": assets,
                "total_amount": total_amount,
                "currency": metadata.get("currency", "CNY"),
                "notes": metadata.get("notes"),
                "raw_data": {
                    "file_name": source_name,
                    "sheet_name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column
                }
            }
            # Keep the flat API used by existing integrations and expose the
            # structured shape consumed by the Streamlit upload page.
            result["extracted_data"] = {
                "project_info": {
                    "project_name": result["project_name"],
                    "client_name": result["client"],
                    "date": result["date"],
                    "contact": result["contact"],
                },
                "assets": result["assets"],
                "total_amount": result["total_amount"],
                "currency": result["currency"],
            }

            logger.info("Excel解析成功")
            return result

        except Exception as e:
            logger.error(f"Excel解析失败: {str(e)}")
            import traceback
            logger.error(f"堆栈跟踪:\n{traceback.format_exc()}")

            return {
                "success": False,
                "error": str(e),
                "error_type": type(e).__name__
            }
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _source_name(source: Union[str, Path, BinaryIO]) -> str:
        if isinstance(source, (str, Path)):
            return Path(source).name
        return Path(getattr(source, "name", "uploaded.xlsx")).name

    @staticmethod
    def _load_workbook(source: Union[str, Path, BinaryIO], source_name: str):
        if hasattr(source, "seek"):
            source.seek(0)

        if Path(source_name).suffix.lower() != ".xls":
            return openpyxl.load_workbook(source, data_only=True)

        # openpyxl does not support legacy .xls files. Convert the first sheet
        # in memory so the same extraction rules can be reused.
        import pandas as pd

        frame = pd.read_excel(source, header=None)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in dataframe_to_rows(frame, index=False, header=False):
            sheet.append([None if pd.isna(value) else value for value in row])
        return workbook

    def _detect_client(self, sheet: Worksheet, hint: str = None) -> str:
        """识别客户"""
        # 如果有提示，优先使用
        if hint:
            for client, template in self.client_templates.items():
                if any(kw in hint for kw in template["keywords"]):
                    return client

        # 扫描前10行寻找客户关键词
        for row in range(1, min(11, sheet.max_row + 1)):
            for col in range(1, min(11, sheet.max_column + 1)):
                cell_value = str(sheet.cell(row, col).value or "").lower()
                for client, template in self.client_templates.items():
                    if any(kw in cell_value for kw in template["keywords"]):
                        return client

        return "未知客户"

    def _extract_metadata(self, sheet: Worksheet, client: str) -> Dict:
        """提取元数据"""
        metadata = {
            "project_name": None,
            "date": None,
            "contact": None,
            "notes": None,
            "currency": "CNY"
        }

        # 扫描前15行寻找关键信息
        for row in range(1, min(16, sheet.max_row + 1)):
            for col in range(1, min(6, sheet.max_column + 1)):
                cell = sheet.cell(row, col)
                value = str(cell.value or "")

                # 项目名称
                if re.search(r'项目名称|project', value, re.I):
                    next_cell = sheet.cell(row, col + 1).value
                    if next_cell:
                        metadata["project_name"] = str(next_cell)

                # 日期
                if re.search(r'日期|date', value, re.I):
                    next_cell = sheet.cell(row, col + 1).value
                    if next_cell:
                        metadata["date"] = self._parse_date(next_cell)

                # 联系人
                if re.search(r'联系人|contact|对接人', value, re.I):
                    next_cell = sheet.cell(row, col + 1).value
                    if next_cell:
                        metadata["contact"] = str(next_cell)

                # 备注
                if re.search(r'备注|说明|note', value, re.I):
                    next_cell = sheet.cell(row, col + 1).value
                    if next_cell:
                        metadata["notes"] = str(next_cell)

        return metadata

    def _find_header_row(self, sheet: Worksheet, template: Dict) -> int:
        """查找表头行"""
        for row in range(1, min(21, sheet.max_row + 1)):
            matched_fields = set()
            for col in range(1, sheet.max_column + 1):
                header = str(sheet.cell(row, col).value or "").strip()
                for field, keywords in template["columns"].items():
                    if any(keyword in header for keyword in keywords):
                        matched_fields.add(field)

            # Metadata rows often contain "项目名称". A real table header must
            # also expose a numeric field such as quantity, unit price or total.
            if "asset_name" in matched_fields and matched_fields.intersection(
                {"quantity", "unit_price", "total"}
            ):
                return row

        return template["header_row"]

    def _map_columns(self, sheet: Worksheet, header_row: int, template: Dict) -> Dict:
        """映射列"""
        column_map = {}

        for col in range(1, sheet.max_column + 1):
            header = str(sheet.cell(header_row, col).value or "").strip()

            for field, keywords in template["columns"].items():
                if any(keyword in header for keyword in keywords):
                    column_map[field] = col
                    break

        return column_map

    def _extract_assets(self, sheet: Worksheet, start_row: int, column_map: Dict) -> List[Dict]:
        """提取资产数据"""
        assets = []

        consecutive_empty_rows = 0
        for row in range(start_row, sheet.max_row + 1):
            # 检查是否是数据行
            asset_name_col = column_map.get("asset_name")
            if not asset_name_col:
                continue

            asset_name = sheet.cell(row, asset_name_col).value
            if not asset_name or not str(asset_name).strip():
                consecutive_empty_rows += 1
                if consecutive_empty_rows >= 3:
                    break
                continue

            consecutive_empty_rows = 0
            if str(asset_name).strip() in ["合计", "总计", "小计"]:
                break

            asset = {
                "name": str(asset_name).strip(),
                "quantity": self._get_number(sheet, row, column_map.get("quantity")),
                "unit_price": self._get_number(sheet, row, column_map.get("unit_price")),
                "total": self._get_number(sheet, row, column_map.get("total"))
            }

            # 计算缺失的值
            if asset["total"] is None and asset["quantity"] and asset["unit_price"]:
                asset["total"] = asset["quantity"] * asset["unit_price"]

            assets.append(asset)

        return assets

    def _get_number(self, sheet: Worksheet, row: int, col: Optional[int]) -> Optional[float]:
        """获取数字值"""
        if col is None:
            return None

        value = sheet.cell(row, col).value
        if value is None:
            return None

        # 尝试转换为数字
        try:
            if isinstance(value, (int, float)):
                return float(value)

            # 移除货币符号和逗号
            value_str = str(value).replace(",", "").replace("¥", "").replace("$", "").strip()
            return float(value_str)
        except (TypeError, ValueError):
            return None

    def _calculate_total(self, assets: List[Dict], sheet: Worksheet) -> float:
        """计算总金额"""
        # Prefer an explicit final total because it may include discounts,
        # taxes or adjustments that are not represented in asset rows.
        for row in range(sheet.max_row, max(0, sheet.max_row - 10), -1):
            for col in range(1, sheet.max_column + 1):
                cell_value = str(sheet.cell(row, col).value or "")
                if re.search(r'合计|总计|总价|total', cell_value, re.I):
                    numeric_values = [
                        self._get_number(sheet, row, candidate)
                        for candidate in range(col + 1, sheet.max_column + 1)
                    ]
                    numeric_values = [value for value in numeric_values if value is not None]
                    if numeric_values:
                        return numeric_values[-1]

        return sum(asset.get("total", 0) or 0 for asset in assets)

    def _parse_date(self, value: Any) -> str:
        """解析日期"""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")

        try:
            date_str = str(value)
            # 尝试多种日期格式
            for fmt in ["%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"]:
                try:
                    return datetime.strptime(date_str, fmt).strftime("%Y-%m-%d")
                except ValueError:
                    continue
            return date_str
        except (TypeError, ValueError, OverflowError):
            return str(value)


# 使用示例
if __name__ == "__main__":
    parser = ExcelQuoteParser()

    # 解析报价单
    result = parser.parse("./test_data/腾讯报价单.xlsx", client_hint="腾讯")

    if result["success"]:
        print(f"客户: {result['client']}")
        print(f"项目: {result['project_name']}")
        print(f"总金额: {result['total_amount']}")
        print("\n资产列表:")
        for asset in result["assets"]:
            print(f"  - {asset['name']}: {asset['quantity']}个 x ¥{asset['unit_price']} = ¥{asset['total']}")
    else:
        print(f"解析失败: {result['error']}")
